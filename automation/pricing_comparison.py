"""Read-only price-snapshot comparison using the partner pack's diff rules."""
from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

try:
    from .partner_export_pack import (
        _read_required_pricing, _currency_view, _shared_filter_rules,
        _clean_partner_table, _partner_diff_table,
    )
    from .plan_labels import partner_display_plan_label
except ImportError:
    from partner_export_pack import (
        _read_required_pricing, _currency_view, _shared_filter_rules,
        _clean_partner_table, _partner_diff_table,
    )
    from plan_labels import partner_display_plan_label

CURRENCIES = ('EUR', 'USD')


def read_snapshot(path):
    """Read an editor workbook or consolidated history CSV; never touch SQLite."""
    path = Path(path)
    if path.suffix.lower() not in {'.xlsx', '.csv'}:
        raise ValueError('Choose a Pricing Editor Excel (.xlsx) or consolidated history CSV.')
    if not path.is_file():
        raise FileNotFoundError(path)
    # Validate raw headers before canonical normalization fills missing columns.
    if path.suffix.lower() == '.csv':
        headers = set(pd.read_csv(path, nrows=0).columns.str.strip())
    else:
        from openpyxl import load_workbook
        book = load_workbook(path, read_only=True, data_only=True)
        try:
            name = 'Pricing' if 'Pricing' in book.sheetnames else ('All_Data' if 'All_Data' in book.sheetnames else book.sheetnames[0])
            headers = {str(value).strip() for value in next(book[name].iter_rows(values_only=True), ()) if value is not None}
        finally:
            book.close()
    required = {'ISO', 'Plan', 'Days', 'GB', 'Price_EUR', 'Price_USD'}
    missing = required - headers
    if missing:
        raise ValueError(f'{path.name}: missing {", ".join(sorted(missing))}. Choose the consolidated editor export/history file, not a partner CSV.')
    frame = _read_required_pricing(path)
    if frame.empty:
        raise ValueError(f'{path.name} contains no pricing rows.')
    return frame


def comparison_tables(previous, current, *, previous_label='Previous', current_label='Current'):
    """Return both full-workbook and partner-eligible final-price changes."""
    snapshots = []
    for frame in (previous, current):
        views = {currency: (_currency_view(frame, currency), frame.iloc[:0].copy()) for currency in CURRENCIES}
        blocked = _shared_filter_rules(views)
        all_prices, partner_prices = {}, {}
        for currency, (prices, regions) in views.items():
            all_prices[currency] = prices.copy()
            all_prices[currency]['Plan'] = all_prices[currency]['Plan'].map(partner_display_plan_label)
            partner_prices[currency], _ = _clean_partner_table(
                prices, regions, currency=currency, shared_blocked_price_keys=blocked,
            )
        snapshots.append((all_prices, partner_prices))
    tables, summary = {}, []
    for index, name in enumerate(('All price changes', 'Partner changes')):
        pieces = []
        for currency in CURRENCIES:
            before, after = snapshots[0][index][currency], snapshots[1][index][currency]
            diff = _partner_diff_table(before, after, currency=currency,
                                       previous_date=previous_label, current_date=current_label)
            pieces.append(diff)
            counts = diff['ChangeType'].value_counts()
            summary.append(dict(View=name, Currency=currency, PreviousProducts=len(before),
                                CurrentProducts=len(after), Changed=int(counts.get('Changed', 0)),
                                Added=int(counts.get('Added', 0)), Removed=int(counts.get('Removed', 0))))
        populated = [piece for piece in pieces if not piece.empty]
        tables[name] = pd.concat(populated, ignore_index=True) if populated else pieces[0].copy()
    return pd.DataFrame(summary), tables


def save_comparison(previous_path, current_path, output_path):
    previous_path, current_path, output_path = map(Path, (previous_path, current_path, output_path))
    if output_path.resolve() in {previous_path.resolve(), current_path.resolve()}:
        raise ValueError('Save the comparison as a separate file; it cannot replace an input.')
    if output_path.suffix.lower() != '.xlsx':
        raise ValueError('Save the comparison as an .xlsx workbook.')
    previous, current = read_snapshot(previous_path), read_snapshot(current_path)
    summary, tables = comparison_tables(previous, current,
        previous_label=previous_path.stem, current_label=current_path.stem)
    notes = pd.DataFrame([
        ('Previous snapshot', str(previous_path.resolve())),
        ('Current snapshot', str(current_path.resolve())),
        ('Comparison', 'Same final-price comparison as partner packs: destination + package + duration + GB, separately for EUR and USD.'),
        ('All price changes', 'Includes every source product, including rows blocked from partner export.'),
        ('Partner changes', 'Uses the same shared EUR/USD cost-floor exclusions and below-cost approvals as partner packs.'),
        ('Changed', 'Final selling price changed. List-price or promo-code changes with identical final prices are not listed.'),
        ('Added / Removed', 'Product entered or left this view; in Partner changes this can also mean export eligibility changed.'),
        ('Unchanged', 'Unchanged products are omitted; an empty change sheet means no differences in final selling prices or product availability.'),
        ('Read only', 'No partner pack, official export, price import, database update or history record is created.'),
    ], columns=['Item', 'Explanation'])
    output_path.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(output_path, engine='openpyxl') as writer:
        summary.to_excel(writer, sheet_name='Summary', index=False)
        for name, table in tables.items():
            table.to_excel(writer, sheet_name=name, index=False)
        notes.to_excel(writer, sheet_name='Notes', index=False)
        for sheet in writer.book.worksheets:
            sheet.freeze_panes = 'A2'
            sheet.auto_filter.ref = sheet.dimensions
            for cell in sheet[1]:
                cell.font = Font(color='FFFFFF', bold=True)
                cell.fill = PatternFill('solid', fgColor='E20074')
            for column in sheet.columns:
                width = max(len(str(cell.value or '')) for cell in column)
                sheet.column_dimensions[get_column_letter(column[0].column)].width = min(48, max(12, width + 2))
            if sheet.title == 'Notes':
                sheet.column_dimensions['B'].width = 100
                for row in sheet.iter_rows(min_row=2):
                    row[1].alignment = Alignment(wrap_text=True, vertical='top')
                    sheet.row_dimensions[row[0].row].height = 42
            elif sheet.title != 'Summary':
                headers = {cell.value: cell.column for cell in sheet[1]}
                for name in ('PreviousPrice', 'CurrentPrice', 'PriceChange',
                             'PreviousFinalPriceAfterPromo', 'CurrentFinalPriceAfterPromo', 'FinalPriceAfterPromoChange'):
                    for column in sheet.iter_cols(min_col=headers[name], max_col=headers[name], min_row=2):
                        for cell in column:
                            cell.number_format = '0.00'
    return summary
