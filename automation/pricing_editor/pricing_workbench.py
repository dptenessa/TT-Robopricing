from __future__ import annotations

from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Any, Iterable
import math

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill, Protection
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.workbook.defined_name import DefinedName

from currency_support import CURRENCIES, normalize_currency
# from automation.currency_support import CURRENCIES, normalize_currency
try:
    from plan_labels import display_plan_label
except ImportError:
    from automation.plan_labels import display_plan_label


WORKBENCH_SHEET = "Pricing Workbench"
PROMO_SHEET = "_PromoCatalog"
META_SHEET = "_Metadata"
WORKBENCH_VERSION = "2"
SUPPORTED_WORKBENCH_VERSIONS = {"1", "2"}
NO_PROMO = "No promo"
YES = "Yes"
NO = "No"

VISIBLE_HEADERS = [
    "Pricing Unit ID",
    "Plan",
    "Days",
    "GB",
    "Country",
    "Countries",
    "List Price EUR",
    "List Price USD",
    "Promo",
    "Promo Value",
    "Net Price EUR",
    "Net Price USD",
    "Cost Floor EUR",
    "Cost Floor USD",
    "Allow Below Cost",
]
HIDDEN_HEADERS = [
    "_ScopeKey",
    "_PromoCode",
    "_PromoType",
    "_PromoRawValue",
]
ALL_HEADERS = VISIBLE_HEADERS + HIDDEN_HEADERS

EDITABLE_HEADERS = {
    "List Price EUR",
    "List Price USD",
    "Promo",
    "Allow Below Cost",
}


@dataclass(frozen=True)
class WorkbenchRow:
    pricing_unit_id: str
    plan: str
    days: float
    gb: float | None
    countries: tuple[str, ...]
    price_eur: float
    price_usd: float
    promo_code: str
    promo_label: str
    country: str
    cost_floor_eur: float
    cost_floor_usd: float
    is_region: bool = False

    @property
    def scope_key(self) -> str:
        return build_scope_key(self.pricing_unit_id, self.plan, self.days, self.gb)


@dataclass(frozen=True)
class WorkbenchImportRow:
    row_number: int
    pricing_unit_id: str
    plan: str
    days: float
    gb: float | None
    countries: str
    price_eur: float
    price_usd: float
    promo_label: str
    promo_code: str
    allow_below_cost: bool
    scope_key: str


@dataclass(frozen=True)
class WorkbenchImportResult:
    rows: list[WorkbenchImportRow]
    pricing_eur_to_usd: float
    generated_at: str
    source_version: str


def _clean_text(value: Any) -> str:
    if value is None:
        return ""
    text = str(value).strip()
    return "" if text.lower() == "nan" else text


def _num_text(value: Any) -> str:
    if value is None:
        return ""
    try:
        number = float(value)
    except (TypeError, ValueError):
        return _clean_text(value)
    if not math.isfinite(number):
        return ""
    if number.is_integer():
        return str(int(number))
    return f"{number:g}"


def build_scope_key(pricing_unit_id: Any, plan: Any, days: Any, gb: Any) -> str:
    return "|".join([
        _clean_text(pricing_unit_id).upper(),
        _clean_text(plan).upper(),
        _num_text(days),
        _num_text(gb),
    ])


def build_partner_price_key(iso: Any, plan: Any, days: Any, gb: Any) -> str:
    return "|".join([
        _clean_text(iso).upper(),
        _clean_text(plan).upper(),
        _num_text(days),
        _num_text(gb),
    ])


def _scope_group_key(point: dict[str, Any]) -> tuple[str, str, float, float | None]:
    unit = _clean_text(point.get("pricing_unit_id"))
    plan = _clean_text(point.get("plan"))
    days = float(point.get("days"))
    gb_value = point.get("gb")
    gb = None if gb_value is None else float(gb_value)
    return unit, plan, days, gb


def _promo_catalog_by_code(state) -> dict[str, dict[str, Any]]:
    return {
        _clean_text(item.get("promo_code")): item
        for item in (state.promo_catalog or [])
        if _clean_text(item.get("promo_code"))
    }


def _promo_catalog_by_label(state) -> dict[str, dict[str, Any]]:
    return {
        _clean_text(item.get("label")) or _clean_text(item.get("promo_code")): item
        for item in (state.promo_catalog or [])
        if _clean_text(item.get("promo_code"))
    }


def workbench_rows_from_state(state) -> list[WorkbenchRow]:
    """Collapse country-level editor rows to one row per commercial pricing scope."""
    groups: dict[tuple[str, str, float, float | None], list[dict[str, Any]]] = {}
    for point in state.row_index.values():
        groups.setdefault(_scope_group_key(point), []).append(point)

    catalog_by_code = _promo_catalog_by_code(state)
    rows: list[WorkbenchRow] = []

    for (unit, plan, days, gb), points in groups.items():
        if not unit or not plan:
            continue

        representative = points[0]
        is_region = _clean_text(representative.get("pricing_source")).lower() == "region_max"
        if is_region:
            countries = tuple(
                state.pricing_unit_country_codes(
                    representative.get("pricing_unit_countries", "")
                )
            )
        else:
            countries = tuple(sorted({
                _clean_text(point.get("iso")).upper()
                for point in points
                if _clean_text(point.get("iso"))
            }))

        prices_by_currency: dict[str, float] = {}
        for currency in CURRENCIES:
            values = {
                round(float(state._working_price_for_currency(point, currency)), 8)
                for point in points
            }
            if len(values) != 1:
                raise ValueError(
                    f"Pricing scope {unit}/{plan}/{_num_text(days)}d has inconsistent "
                    f"{currency} list prices across countries: {sorted(values)}"
                )
            prices_by_currency[currency] = float(next(iter(values)))

        promo_key = _clean_text(representative.get("promo_scope_key"))
        promo_codes: dict[str, str] = {}
        for currency in CURRENCIES:
            promo = state._promo_store_for(currency).get(promo_key) if promo_key else None
            promo_codes[currency] = _clean_text((promo or {}).get("promo_code"))

        if len(set(promo_codes.values())) > 1:
            raise ValueError(
                f"Pricing scope {unit}/{plan}/{_num_text(days)}d has different promo assignments "
                f"between EUR and USD ({promo_codes}). Synchronize them before exporting a workbench."
            )

        promo_code = promo_codes.get("EUR") or promo_codes.get("USD") or ""
        promo_item = catalog_by_code.get(promo_code, {})
        promo_label = _clean_text(promo_item.get("label")) or promo_code
        
        state._refresh_point_floor_status(representative)
        floors = representative.get("cost_floor_by_currency", {}) or {}
        cost_floor_eur = floors.get("EUR", "")
        cost_floor_usd = floors.get("USD", "")

        rows.append(WorkbenchRow(
            pricing_unit_id=unit,
            plan=plan,
            days=days,
            gb=gb,
            countries=countries,
            price_eur=prices_by_currency["EUR"],
            price_usd=prices_by_currency["USD"],
            promo_code=promo_code,
            promo_label=promo_label,
            country=_clean_text(representative.get("country")),
            cost_floor_eur=float(cost_floor_eur) if cost_floor_eur != "" else 0.0,
            cost_floor_usd=float(cost_floor_usd) if cost_floor_usd != "" else 0.0,
            is_region=is_region,
        ))

    plan_order = {"basic": 0, "medium": 1, "moderate": 1, "large": 2, "unlimited": 3}
    rows.sort(key=lambda row: (
        1 if row.is_region else 0,
        row.pricing_unit_id.upper(),
        plan_order.get(row.plan.lower(), 99),
        row.plan.lower(),
        row.days,
        -1 if row.gb is None else row.gb,
    ))
    return rows


def _set_defined_name(workbook: Workbook, name: str, attr_text: str) -> None:
    defined_name = DefinedName(name, attr_text=attr_text)
    try:
        workbook.defined_names.add(defined_name)
    except AttributeError:
        workbook.defined_names.append(defined_name)


def export_pricing_workbench(state, path: str | Path) -> Path:
    path = Path(path)
    if path.suffix.lower() != ".xlsx":
        path = path.with_suffix(".xlsx")
    path.parent.mkdir(parents=True, exist_ok=True)

    rows = workbench_rows_from_state(state)
    if not rows:
        raise ValueError("No pricing rows are loaded in the editor.")

    wb = Workbook()
    try:
        wb.calculation.fullCalcOnLoad = True
        wb.calculation.forceFullCalc = True
        wb.calculation.calcMode = "auto"
    except Exception:
        pass
    ws = wb.active
    ws.title = WORKBENCH_SHEET
    promo_ws = wb.create_sheet(PROMO_SHEET)
    meta_ws = wb.create_sheet(META_SHEET)

    # Promo catalogue. Keep 'No promo' as an explicit selectable item.
    promo_ws.append(["Promo", "PromoCode", "PromoType", "PromoValue"])
    promo_ws.append([NO_PROMO, "", "", ""])
    for promo in state.promo_catalog or []:
        code = _clean_text(promo.get("promo_code"))
        if not code:
            continue
        label = _clean_text(promo.get("label")) or code
        promo_type = _clean_text(promo.get("promo_type")).lower()
        promo_value = float(promo.get("promo_value", 0) or 0)
        promo_ws.append([label, code, promo_type, promo_value])

    promo_count = promo_ws.max_row
    _set_defined_name(wb, "PromoChoices", f"'{PROMO_SHEET}'!$A$2:$A${promo_count}")

    # Metadata used by formulas and import validation.
    meta_ws.append(["Key", "Value"])
    metadata = [
        ("workbench_version", WORKBENCH_VERSION),
        ("generated_at", datetime.now().isoformat(timespec="seconds")),
        ("pricing_eur_to_usd", float(state.eur_to_usd)),
        ("row_count", len(rows)),
    ]
    for key, value in metadata:
        meta_ws.append([key, value])

    ws.append(ALL_HEADERS)
    for item in rows:
        ws.append([
            item.pricing_unit_id,
            item.plan,
            item.days,
            item.gb if item.gb is not None else "",
            item.country,
            ", ".join(item.countries),
            item.price_eur,
            item.price_usd,
            item.promo_label or NO_PROMO,
            "",  # Promo Value (formula written below)
            "",  # Net Price EUR (formula written below)
            "",  # Net Price USD (formula written below)
            item.cost_floor_eur,
            item.cost_floor_usd,
            NO,
            item.scope_key,
            "", "", "",
        ])

    header_fill = PatternFill("solid", fgColor="30343B")
    header_font = Font(color="FFFFFF", bold=True)
    editable_fill = PatternFill("solid", fgColor="FFF2CC")
    formula_fill = PatternFill("solid", fgColor="E2F0D9")
    warning_fill = PatternFill("solid", fgColor="FCE4D6")

    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.protection = Protection(locked=True)

    header_index = {cell.value: cell.column for cell in ws[1]}
    for row_idx in range(2, ws.max_row + 1):
        promo_col = get_column_letter(header_index["Promo"])
        list_eur_col = get_column_letter(header_index["List Price EUR"])
        list_usd_col = get_column_letter(header_index["List Price USD"])
        promo_value_col = get_column_letter(header_index["Promo Value"])
        net_eur_col = get_column_letter(header_index["Net Price EUR"])
        net_usd_col = get_column_letter(header_index["Net Price USD"])
        helper_code_col = get_column_letter(header_index["_PromoCode"])
        helper_type_col = get_column_letter(header_index["_PromoType"])
        helper_value_col = get_column_letter(header_index["_PromoRawValue"])

        ws[f"{helper_code_col}{row_idx}"] = f"=IFERROR(VLOOKUP(${promo_col}{row_idx},'{PROMO_SHEET}'!$A:$D,2,FALSE),\"\")"
        ws[f"{helper_type_col}{row_idx}"] = f"=IFERROR(VLOOKUP(${promo_col}{row_idx},'{PROMO_SHEET}'!$A:$D,3,FALSE),\"\")"
        ws[f"{helper_value_col}{row_idx}"] = f"=IFERROR(VLOOKUP(${promo_col}{row_idx},'{PROMO_SHEET}'!$A:$D,4,FALSE),\"\")"

        ws[f"{promo_value_col}{row_idx}"] = (
            f'=IF(${promo_col}{row_idx}="{NO_PROMO}","",'
            f'IF(${helper_type_col}{row_idx}="percent",TEXT(${helper_value_col}{row_idx}/100,"0%"),'
            f'TEXT(${helper_value_col}{row_idx},"0.00")))'
        )

        eur_discount = (
            f'IF(${helper_type_col}{row_idx}="percent",${list_eur_col}{row_idx}*${helper_value_col}{row_idx}/100,'
            f'${helper_value_col}{row_idx})'
        )
        usd_discount = (
            f'IF(${helper_type_col}{row_idx}="percent",${list_usd_col}{row_idx}*${helper_value_col}{row_idx}/100,'
            f'${helper_value_col}{row_idx})'
        )
        ws[f"{net_eur_col}{row_idx}"] = (
            f'=IF(${promo_col}{row_idx}="{NO_PROMO}",${list_eur_col}{row_idx},'
            f'ROUNDDOWN(MAX(0,${list_eur_col}{row_idx}-({eur_discount}))*20,0)/20)'
        )
        ws[f"{net_usd_col}{row_idx}"] = (
            f'=IF(${promo_col}{row_idx}="{NO_PROMO}",${list_usd_col}{row_idx},'
            f'ROUNDDOWN(MAX(0,${list_usd_col}{row_idx}-({usd_discount}))*20,0)/20)'
        )

        for header in ALL_HEADERS:
            cell = ws.cell(row=row_idx, column=header_index[header])
            cell.protection = Protection(locked=header not in EDITABLE_HEADERS)
            cell.alignment = Alignment(vertical="center")

        for header in EDITABLE_HEADERS:
            ws.cell(row=row_idx, column=header_index[header]).fill = editable_fill
        for header in ("Promo Value", "Net Price EUR", "Net Price USD"):
            ws.cell(row=row_idx, column=header_index[header]).fill = formula_fill

    # Dropdowns.
    promo_dv = DataValidation(type="list", formula1="=PromoChoices", allow_blank=False)
    promo_dv.error = "Select a promo from the approved promo catalogue."
    promo_dv.errorTitle = "Invalid promo"
    promo_dv.prompt = "Choose an approved promo."
    promo_dv.promptTitle = "Promo"
    ws.add_data_validation(promo_dv)
    promo_letter = get_column_letter(header_index["Promo"])
    promo_dv.add(f"{promo_letter}2:{promo_letter}{ws.max_row}")

    override_dv = DataValidation(type="list", formula1='"No,Yes"', allow_blank=False)
    override_dv.error = "Choose Yes or No."
    override_dv.errorTitle = "Invalid value"
    ws.add_data_validation(override_dv)
    override_letter = get_column_letter(header_index["Allow Below Cost"])
    override_dv.add(f"{override_letter}2:{override_letter}{ws.max_row}")

    # Workbook usability.
    ws.freeze_panes = "A2"
    last_visible_col = get_column_letter(len(VISIBLE_HEADERS))
    ws.auto_filter.ref = f"A1:{last_visible_col}{ws.max_row}"
    widths_by_header = {
        "Pricing Unit ID": 18,
        "Plan": 14,
        "Days": 9,
        "GB": 10,
        "Country": 24,
        "Countries": 28,
        "List Price EUR": 15,
        "List Price USD": 15,
        "Promo": 20,
        "Promo Value": 14,
        "Net Price EUR": 15,
        "Net Price USD": 15,
        "Cost Floor EUR": 15,
        "Cost Floor USD": 15,
        "Allow Below Cost": 18,
    }
    for header, width in widths_by_header.items():
        ws.column_dimensions[get_column_letter(header_index[header])].width = width

    for col in range(len(VISIBLE_HEADERS) + 1, len(ALL_HEADERS) + 1):
        ws.column_dimensions[get_column_letter(col)].hidden = True

    for header in (
        "List Price EUR", "List Price USD",
        "Net Price EUR", "Net Price USD",
        "Cost Floor EUR", "Cost Floor USD",
    ):
        col_idx = header_index[header]
        for row_idx in range(2, ws.max_row + 1):
            ws.cell(row=row_idx, column=col_idx).number_format = '0.00'

    # Make deliberate below-cost overrides visually obvious.
    ws.conditional_formatting.add(
        f"{override_letter}2:{override_letter}{ws.max_row}",
        FormulaRule(formula=[f'{override_letter}2="Yes"'], fill=warning_fill),
    )

    promo_ws.sheet_state = "hidden"
    meta_ws.sheet_state = "hidden"
    wb.save(path)
    return path


def _metadata_map(workbook) -> dict[str, Any]:
    if META_SHEET not in workbook.sheetnames:
        raise ValueError(f"Workbook is missing required sheet {META_SHEET!r}.")
    ws = workbook[META_SHEET]
    return {
        _clean_text(ws.cell(row=row, column=1).value): ws.cell(row=row, column=2).value
        for row in range(2, ws.max_row + 1)
        if _clean_text(ws.cell(row=row, column=1).value)
    }


def _promo_lookup_from_workbook(workbook) -> dict[str, dict[str, Any]]:
    if PROMO_SHEET not in workbook.sheetnames:
        raise ValueError(f"Workbook is missing required sheet {PROMO_SHEET!r}.")
    ws = workbook[PROMO_SHEET]
    lookup: dict[str, dict[str, Any]] = {}
    for row in range(2, ws.max_row + 1):
        label = _clean_text(ws.cell(row=row, column=1).value)
        if not label:
            continue
        lookup[label] = {
            "promo_code": _clean_text(ws.cell(row=row, column=2).value),
            "promo_type": _clean_text(ws.cell(row=row, column=3).value).lower(),
            "promo_value": ws.cell(row=row, column=4).value,
        }
    return lookup


def read_pricing_workbench(path: str | Path) -> WorkbenchImportResult:
    path = Path(path)
    if not path.exists():
        raise FileNotFoundError(path)

    wb = load_workbook(path, data_only=False, read_only=False)
    if WORKBENCH_SHEET not in wb.sheetnames:
        raise ValueError(f"Workbook is missing required sheet {WORKBENCH_SHEET!r}.")

    metadata = _metadata_map(wb)
    version = _clean_text(metadata.get("workbench_version"))
    if version not in SUPPORTED_WORKBENCH_VERSIONS:
        raise ValueError(
            f"Unsupported pricing workbench version {version!r}; supported versions are "
            + ", ".join(sorted(SUPPORTED_WORKBENCH_VERSIONS))
            + "."
        )
    try:
        rate = float(metadata.get("pricing_eur_to_usd"))
    except (TypeError, ValueError):
        raise ValueError("Workbook metadata contains an invalid pricing EUR/USD rate.")
    if rate <= 0:
        raise ValueError("Workbook metadata contains a non-positive pricing EUR/USD rate.")

    promo_lookup = _promo_lookup_from_workbook(wb)
    ws = wb[WORKBENCH_SHEET]
    header_map = {
        _clean_text(ws.cell(row=1, column=col).value): col
        for col in range(1, ws.max_column + 1)
    }
    missing = [header for header in ALL_HEADERS if header not in header_map]
    if missing:
        raise ValueError("Workbook is missing required columns: " + ", ".join(missing))

    rows: list[WorkbenchImportRow] = []
    seen_scope_keys: set[str] = set()
    for row_idx in range(2, ws.max_row + 1):
        unit = _clean_text(ws.cell(row=row_idx, column=header_map["Pricing Unit ID"]).value)
        plan = _clean_text(ws.cell(row=row_idx, column=header_map["Plan"]).value)
        if not unit and not plan:
            continue

        try:
            days = float(ws.cell(row=row_idx, column=header_map["Days"]).value)
        except (TypeError, ValueError):
            raise ValueError(f"Row {row_idx}: Days must be numeric.")

        gb_raw = ws.cell(row=row_idx, column=header_map["GB"]).value
        if gb_raw in (None, ""):
            gb = None
        else:
            try:
                gb = float(gb_raw)
            except (TypeError, ValueError):
                raise ValueError(f"Row {row_idx}: GB must be numeric or blank.")

        def price(header: str) -> float:
            raw = ws.cell(row=row_idx, column=header_map[header]).value
            try:
                value = float(raw)
            except (TypeError, ValueError):
                raise ValueError(f"Row {row_idx}: {header} must be numeric.")
            if not math.isfinite(value) or value < 0:
                raise ValueError(f"Row {row_idx}: {header} must be a non-negative finite number.")
            rounded = round(value * 20) / 20
            if abs(value - rounded) > 1e-8:
                raise ValueError(
                    f"Row {row_idx}: {header}={value:g} is not on the required 0.05 price grid."
                )
            return rounded

        price_eur = price("List Price EUR")
        price_usd = price("List Price USD")
        countries = _clean_text(ws.cell(row=row_idx, column=header_map["Countries"]).value)
        promo_label = _clean_text(ws.cell(row=row_idx, column=header_map["Promo"]).value) or NO_PROMO
        if promo_label not in promo_lookup:
            raise ValueError(
                f"Row {row_idx}: promo {promo_label!r} is not in the approved promo catalogue embedded in the workbook."
            )
        promo_code = _clean_text(promo_lookup[promo_label].get("promo_code"))

        override_text = _clean_text(ws.cell(row=row_idx, column=header_map["Allow Below Cost"]).value) or NO
        if override_text.lower() not in {YES.lower(), NO.lower()}:
            raise ValueError(f"Row {row_idx}: Allow Below Cost must be Yes or No.")
        allow_below = override_text.lower() == YES.lower()

        actual_scope_key = build_scope_key(unit, plan, days, gb)
        embedded_scope_key = _clean_text(ws.cell(row=row_idx, column=header_map["_ScopeKey"]).value)
        if embedded_scope_key != actual_scope_key:
            raise ValueError(
                f"Row {row_idx}: structural fields were changed. Expected scope key "
                f"{embedded_scope_key!r}, now {actual_scope_key!r}. Only prices, promo and below-cost override may be edited."
            )
        if actual_scope_key in seen_scope_keys:
            raise ValueError(f"Row {row_idx}: duplicate pricing scope {actual_scope_key}.")
        seen_scope_keys.add(actual_scope_key)

        rows.append(WorkbenchImportRow(
            row_number=row_idx,
            pricing_unit_id=unit,
            plan=plan,
            days=days,
            gb=gb,
            countries=countries,
            price_eur=price_eur,
            price_usd=price_usd,
            promo_label=promo_label,
            promo_code=promo_code,
            allow_below_cost=allow_below,
            scope_key=actual_scope_key,
        ))

    expected_count = int(metadata.get("row_count") or 0)
    if expected_count and len(rows) != expected_count:
        raise ValueError(
            f"Workbook row count changed ({len(rows)} rows found; expected {expected_count}). "
            "Do not add or delete pricing rows."
        )

    return WorkbenchImportResult(
        rows=rows,
        pricing_eur_to_usd=rate,
        generated_at=_clean_text(metadata.get("generated_at")),
        source_version=version,
    )


def validate_workbench_against_state(imported: WorkbenchImportResult, state) -> None:
    current = {row.scope_key: row for row in workbench_rows_from_state(state)}
    incoming = {row.scope_key: row for row in imported.rows}
    missing = sorted(set(current) - set(incoming))
    extra = sorted(set(incoming) - set(current))
    if missing or extra:
        raise ValueError(
            "Workbook structure no longer matches the current pricing model. "
            f"Missing scopes: {missing[:10]}; extra scopes: {extra[:10]}. "
            "Generate a fresh workbench from the GUI."
        )

    # Countries are display-only but are also a useful stale-workbook guard.
    for key, incoming_row in incoming.items():
        expected_countries = ", ".join(current[key].countries)
        if _clean_text(incoming_row.countries) != expected_countries:
            raise ValueError(
                f"Workbook scope {key} has a changed/stale Countries field. "
                "Generate a fresh workbench instead of editing structural columns."
            )


def apply_workbench_to_state(imported: WorkbenchImportResult, state) -> dict[str, Any]:
    """Apply list-price and promo decisions to EditorState and return change/override details."""
    validate_workbench_against_state(imported, state)
    state.set_eur_to_usd(imported.pricing_eur_to_usd)

    current_rows = {row.scope_key: row for row in workbench_rows_from_state(state)}
    price_changed_eur = 0
    price_changed_usd = 0
    promo_changed = 0
    changed_scopes = 0
    override_scope_keys: set[str] = set()
    override_partner_keys: set[str] = set()
    changed_scope_keys: set[str] = set()

    for item in imported.rows:
        before = current_rows[item.scope_key]
        price_eur_changed = abs(float(before.price_eur) - float(item.price_eur)) > 1e-9
        price_usd_changed = abs(float(before.price_usd) - float(item.price_usd)) > 1e-9
        before_promo = _clean_text(before.promo_code)
        promo_changed_here = before_promo != _clean_text(item.promo_code)

        matches = state.workbench_scope_points(
            item.pricing_unit_id,
            item.plan,
            item.days,
            item.gb,
        )
        if not matches:
            raise ValueError(f"Row {item.row_number}: no current pricing points match {item.scope_key}.")

        state.set_workbench_scope_prices(
            item.pricing_unit_id,
            item.plan,
            item.days,
            item.gb,
            price_eur=item.price_eur,
            price_usd=item.price_usd,
        )
        state.set_workbench_scope_promo(
            item.pricing_unit_id,
            item.plan,
            item.days,
            item.gb,
            promo_code=item.promo_code,
        )

        if price_eur_changed:
            price_changed_eur += 1
        if price_usd_changed:
            price_changed_usd += 1
        if promo_changed_here:
            promo_changed += 1
        if price_eur_changed or price_usd_changed or promo_changed_here or item.allow_below_cost:
            changed_scopes += 1
            changed_scope_keys.add(item.scope_key)

        if item.allow_below_cost:
            override_scope_keys.add(item.scope_key)
            for point in matches:
                override_partner_keys.add(build_partner_price_key(
                    point.get("iso", ""),
                    point.get("plan", ""),
                    point.get("days", ""),
                    point.get("gb", ""),
                ))

    # Refresh floors/net prices after all changes.
    state._refresh_all_display_prices()

    below_total = 0
    below_overridden = 0
    below_not_overridden = 0
    for point in state.row_index.values():
        below = bool(point.get("is_partner_export_blocked"))
        if not below:
            continue
        below_total += 1
        key = build_partner_price_key(
            point.get("iso", ""), point.get("plan", ""), point.get("days", ""), point.get("gb", "")
        )
        if key in override_partner_keys:
            below_overridden += 1
        else:
            below_not_overridden += 1

    return {
        "changed_scopes": changed_scopes,
        "price_changed_eur": price_changed_eur,
        "price_changed_usd": price_changed_usd,
        "promo_changed": promo_changed,
        "override_scope_count": len(override_scope_keys),
        "override_partner_keys": override_partner_keys,
        "below_total": below_total,
        "below_overridden": below_overridden,
        "below_not_overridden": below_not_overridden,
        "changed_scope_keys": changed_scope_keys,
    }
