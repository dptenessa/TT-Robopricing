from __future__ import annotations
import math
import json
import os
import shutil
import tempfile
import time
from dataclasses import dataclass
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Callable

from openpyxl import load_workbook
import pandas as pd

from .store import Campaign

PRICING_SHEET = "Pricing"
PROMOS_SHEET = "Promos"
PROMO_NAME_COLUMN = "PromoName"
ORIGINAL_PRICE_EUR_COLUMN = "_PromoPlannerOriginalPrice_EUR"
ORIGINAL_PRICE_USD_COLUMN = "_PromoPlannerOriginalPrice_USD"
PRICE_OWNER_COLUMN = "_PromoPlannerPriceOwner"
PROMO_STACK_COLUMN = "_PromoPlannerPromoStack"
PROMO_OWNER_COLUMN = "_PromoPlannerPromoOwner"


@dataclass
class PromoOption:
    code: str
    promo_type: str
    value: float
    label: str

    @property
    def display(self) -> str:
        suffix = "%" if self.promo_type == "percent" else ""
        return f"{self.label}  [{self.code}]  · {self.value:g}{suffix}"


@dataclass
class Catalogue:
    destinations: list[tuple[str, str]]
    plans: list[str]
    days: list[int]
    promos: list[PromoOption]
    row_index: dict[tuple[str, str, int], list[int]]
    workbook_signature: tuple[int, int]


@dataclass
class Preview:
    action: str
    matched_rows: int
    changed_rows: int
    below_floor_rows: int
    collisions: list[str]
    examples: list[str]


def _text(v) -> str:
    if v is None:
        return ""
    s = str(v).strip()
    return "" if s.lower() in {"nan", "none"} else s


def _num(v):
    try:
        if v is None or (isinstance(v, str) and v.startswith("=")):
            return None
        value = float(v)
        return value if math.isfinite(value) else None
    except Exception:
        return None


def round_promo_price(base: float, ptype: str, value: float) -> float:
    if ptype == "percent":
        raw = base * (1.0 - value / 100.0)
    elif ptype == "absolute":
        raw = base - value
    elif ptype == "net":
        raw = value
    else:
        raw = base
    return math.floor(max(raw, 0.0) * 20.0 + 1e-12) / 20.0


def _headers(ws) -> dict[str, int]:
    return {_text(cell.value): cell.column for cell in ws[1] if _text(cell.value)}


def _find_col(h: dict[str, int], *names: str) -> int | None:
    lower = {k.lower(): v for k, v in h.items()}
    for name in names:
        if name.lower() in lower:
            return lower[name.lower()]
    return None


def _read_promos(wb) -> list[PromoOption]:
    if PROMOS_SHEET not in wb.sheetnames:
        return []
    ws = wb[PROMOS_SHEET]
    h = _headers(ws)
    c_code = _find_col(h, "PromoCode", "Code", "_PromoCode")
    c_type = _find_col(h, "PromoType", "Type", "_Type")
    c_value = _find_col(h, "PromoValue", "Value", "RawValue", "_RawValue")
    c_label = _find_col(h, "PromoLabel", "PromoName", "Name", "Label")
    if not (c_code and c_type and c_value):
        return []
    out=[]
    for r in range(2, ws.max_row+1):
        code=_text(ws.cell(r,c_code).value)
        ptype=_text(ws.cell(r,c_type).value).lower()
        value=_num(ws.cell(r,c_value).value)
        label=_text(ws.cell(r,c_label).value) if c_label else code
        if not code or value is None:
            continue
        if ptype in {"percentage","pct","%"}: ptype="percent"
        if ptype in {"fixed","amount","value"}: ptype="absolute"
        if ptype not in {"percent","absolute"}: continue
        out.append(PromoOption(code,ptype,float(value),label or code))
    return sorted(out,key=lambda x:(x.label.lower(),x.code.lower()))


def read_catalogue(path: Path) -> Catalogue:
    wb = load_workbook(path, read_only=True, data_only=False)
    try:
        ws = wb[PRICING_SHEET] if PRICING_SHEET in wb.sheetnames else wb[wb.sheetnames[0]]
        h = _headers(ws)
        required = {"Country", "ISO", "Plan", "Days"}
        missing = required - set(h)
        if missing:
            raise ValueError("Pricing workbook is missing columns: " + ", ".join(sorted(missing)))
        dest: dict[str, str] = {}
        plans = set(); days = set()
        row_index: dict[tuple[str, str, int], list[int]] = {}
        for row_number, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            code = _text(row[h["ISO"] - 1]).upper()
            country = _text(row[h["Country"] - 1])
            plan = _text(row[h["Plan"] - 1])
            day = _num(row[h["Days"] - 1])
            if code: dest[code] = country or code
            if plan: plans.add(plan)
            if day is not None and day > 0:
                day_int = int(day)
                days.add(day_int)
                if code and plan:
                    row_index.setdefault((code, plan, day_int), []).append(row_number)
        stat = Path(path).stat()
        return Catalogue(
            sorted(dest.items(), key=lambda x: (x[1].lower(), x[0])),
            sorted(plans), sorted(days), _read_promos(wb), row_index,
            (stat.st_mtime_ns, stat.st_size),
        )
    finally:
        wb.close()


def read_catalogue_frame(frame: pd.DataFrame, promos: list[PromoOption] | None = None) -> Catalogue:
    """Build the planner catalogue from the shared pricing projection."""
    destinations: dict[str, str] = {}
    plans: set[str] = set()
    days: set[int] = set()
    row_index: dict[tuple[str, str, int], list[int]] = {}
    for row_number, row in enumerate(frame.to_dict(orient="records"), start=2):
        code = _text(row.get("ISO")).upper()
        country = _text(row.get("Country"))
        plan = _text(row.get("Plan"))
        day = _num(row.get("Days"))
        if code:
            destinations.setdefault(code, country or code)
        if plan:
            plans.add(plan)
        if day is not None and day > 0:
            day_int = int(day)
            days.add(day_int)
            if code and plan:
                row_index.setdefault((code, plan, day_int), []).append(row_number)
    if promos is None:
        found: dict[str, PromoOption] = {}
        for row in frame.to_dict(orient="records"):
            code = _text(row.get("PromoCode"))
            ptype = _text(row.get("PromoType")).lower()
            value = _num(row.get("PromoValue"))
            if not code or value is None:
                continue
            if ptype in {"percentage", "pct", "%"}:
                ptype = "percent"
            if ptype in {"fixed", "amount", "value"}:
                ptype = "absolute"
            if ptype in {"percent", "absolute"}:
                found[code] = PromoOption(
                    code, ptype, float(value), _text(row.get("PromoLabel")) or code
                )
        promos = list(found.values())
    return Catalogue(
        sorted(destinations.items(), key=lambda item: (item[1].casefold(), item[0])),
        sorted(plans), sorted(days), sorted(promos, key=lambda item: (item.label.casefold(), item.code.casefold())),
        row_index, (0, 0),
    )


def _matches(row, h: dict[str, int], c: Campaign) -> bool:
    iso = _text(row[h["ISO"] - 1]).upper()
    plan = _text(row[h["Plan"] - 1])
    day = _num(row[h["Days"] - 1])
    return iso in set(c.destinations) and plan in set(c.plans) and day is not None and int(day) in set(c.days)


def preview_campaign(path: Path | pd.DataFrame, c: Campaign, action: str) -> Preview:
    action = action.upper()
    is_frame = isinstance(path, pd.DataFrame)
    wb = None if is_frame else load_workbook(path, read_only=True, data_only=False)
    try:
        if is_frame:
            frame = path.astype(object).where(pd.notna(path), None).reset_index(drop=True)
            h = {str(column): index + 1 for index, column in enumerate(frame.columns)}
            rows = frame.itertuples(index=False, name=None)
            promo_catalogue = {}
        else:
            ws = wb[PRICING_SHEET] if PRICING_SHEET in wb.sheetnames else wb[wb.sheetnames[0]]
            h = _headers(ws)
            rows = ws.iter_rows(min_row=2, values_only=True)
        for col in ("ISO", "Plan", "Days", "PromoCode"):
            if col not in h: raise ValueError(f"Pricing workbook is missing {col}.")
        matched = changed = below = 0; collisions=[]; examples=[]
        for values in rows:
            if not _matches(values, h, c): continue
            matched += 1
            current = _text(values[h["PromoCode"] - 1])
            iso = _text(values[h["ISO"] - 1]); plan = _text(values[h["Plan"] - 1]); day = int(_num(values[h["Days"] - 1]) or 0)
            if action == "START":
                if current and current != c.promo_code:
                    collisions.append(f"{iso} / {plan} / {day}d already has {current}"); continue
                if current != c.promo_code: changed += 1
                for currency in ("EUR", "USD"):
                    pcol, fcol = f"Price_{currency}", f"CostFloor_{currency}"
                    if pcol in h and fcol in h:
                        base = _num(values[h[pcol]-1]); floor = _num(values[h[fcol]-1])
                        if base is not None and floor is not None and round_promo_price(base,c.promo_type,c.promo_value) < floor - 1e-9:
                            below += 1; break
            elif action == "END" and current == c.promo_code:
                changed += 1
            if len(examples) < 8: examples.append(f"{iso} · {plan} · {day}d · current promo: {current or 'none'}")
        return Preview(action, matched, changed, below, collisions[:100], examples)
    finally:
        if wb is not None:
            wb.close()


def _ensure_promo_name_column(ws, h: dict[str,int]) -> dict[str,int]:
    if PROMO_NAME_COLUMN not in h:
        col = ws.max_column + 1
        ws.cell(1,col).value = PROMO_NAME_COLUMN
        h[PROMO_NAME_COLUMN]=col
    return h


def _ensure_planner_price_columns(ws, h: dict[str, int]) -> dict[str, int]:
    """Ensure hidden columns used to restore temporary list-price uplifts."""
    for name in (
        ORIGINAL_PRICE_EUR_COLUMN, ORIGINAL_PRICE_USD_COLUMN, PRICE_OWNER_COLUMN,
        PROMO_STACK_COLUMN, PROMO_OWNER_COLUMN,
    ):
        if name not in h:
            col = ws.max_column + 1
            ws.cell(1, col).value = name
            h[name] = col
        try:
            ws.column_dimensions[ws.cell(1, h[name]).column_letter].hidden = True
        except Exception:
            pass
    return h


def execute_campaign(path: Path, c: Campaign, action: str, backup_dir: Path) -> tuple[Path, int]:
    action = action.upper()
    preview = preview_campaign(path,c,action)
    if action == "START" and preview.collisions:
        raise RuntimeError("Promo collision detected. Resolve before execution:\n" + "\n".join(preview.collisions[:20]))
    if action == "START" and preview.below_floor_rows and not c.allow_below_cost:
        raise RuntimeError(f"{preview.below_floor_rows} rows would be below cost floor. Enable 'Allow below cost' explicitly or change the promotion.")
    backup_dir.mkdir(parents=True,exist_ok=True)
    ts=datetime.now().strftime("%Y%m%d-%H%M%S")
    backup=backup_dir / f"{path.stem}_before_{c.promo_code}_{action}_{ts}{path.suffix}"
    shutil.copy2(path,backup)
    wb=load_workbook(path,keep_vba=path.suffix.lower()==".xlsm")
    try:
        ws=wb[PRICING_SHEET] if PRICING_SHEET in wb.sheetnames else wb[wb.sheetnames[0]]
        h=_ensure_promo_name_column(ws,_headers(ws))
        changed=0
        for r in range(2,ws.max_row+1):
            vals=[ws.cell(r,col).value for col in range(1,ws.max_column+1)]
            if not _matches(vals,h,c): continue
            promo_cell=ws.cell(r,h["PromoCode"]); name_cell=ws.cell(r,h[PROMO_NAME_COLUMN]); current=_text(promo_cell.value)
            if action=="START":
                if current==c.promo_code:
                    if _text(name_cell.value)!=c.promo_name: name_cell.value=c.promo_name
                    continue
                promo_cell.value=c.promo_code; name_cell.value=c.promo_name
                if "AllowBelowCost" in h: ws.cell(r,h["AllowBelowCost"]).value=False
                changed += 1
            elif action=="END" and current==c.promo_code:
                promo_cell.value=""; name_cell.value=""; changed += 1
        try:
            wb.calculation.fullCalcOnLoad=True; wb.calculation.forceFullCalc=True; wb.calculation.calcMode="auto"
        except Exception: pass
        wb.save(path)
    except Exception:
        try: wb.close()
        except Exception: pass
        shutil.copy2(backup,path); raise
    finally:
        try: wb.close()
        except Exception: pass
    return backup,changed


@dataclass
class BatchPreviewItem:
    campaign_id: int
    campaign_name: str
    action: str
    matched_rows: int
    changed_rows: int
    below_floor_rows: int
    collisions: list[str]


@dataclass
class CostFloorIssue:
    campaign_id: int
    campaign_name: str
    iso: str
    plan: str
    days: int
    currency: str
    current_list_price: float
    promo_net_price: float
    cost_floor: float
    target_net_price: float
    required_list_price: float
    is_adjustment: bool = False


@dataclass
class PromoCollisionIssue:
    campaign_id: int
    campaign_name: str
    iso: str
    plan: str
    days: int
    current_promo_code: str
    current_promo_name: str


@dataclass
class PriceChange:
    iso: str
    country: str
    plan: str
    days: int
    change: str
    promo_before: str
    promo_after: str
    list_eur_before: float | None
    list_eur_after: float | None
    net_eur_before: float | None
    net_eur_after: float | None
    list_usd_before: float | None
    list_usd_after: float | None
    net_usd_before: float | None
    net_usd_after: float | None
    allow_below_cost_before: bool
    allow_below_cost_after: bool


@dataclass
class BatchPreview:
    execution_date: str
    items: list[BatchPreviewItem]
    cost_floor_issues: list[CostFloorIssue]
    collision_issues: list[PromoCollisionIssue]
    price_changes: list[PriceChange]
    projected_rows: list[dict]

    @property
    def total_changed_rows(self) -> int:
        return sum(i.changed_rows for i in self.items)

    @property
    def collisions(self) -> list[str]:
        out: list[str] = []
        for i in self.items:
            out.extend(i.collisions)
        return out

    @property
    def below_floor_rows(self) -> int:
        return sum(i.below_floor_rows for i in self.items)


def _ceil_price_grid(value: float) -> float:
    return math.ceil((float(value) - 1e-12) * 20.0) / 20.0


def required_list_price_for_target(target_net: float, ptype: str, value: float) -> float:
    """Return the smallest 0.05-grid list price whose promo net is >= target_net."""
    if ptype == "percent":
        factor = 1.0 - float(value) / 100.0
        if factor <= 0:
            raise ValueError("A 100% (or greater) discount cannot be lifted above the cost floor by changing list price.")
        raw = float(target_net) / factor
    elif ptype == "absolute":
        raw = float(target_net) + float(value)
    else:
        raw = float(target_net)
    candidate = _ceil_price_grid(raw)
    # Promo calculation rounds DOWN to the 0.05 grid, so verify and nudge if needed.
    while round_promo_price(candidate, ptype, value) < float(target_net) - 1e-9:
        candidate = round(candidate + 0.05, 10)
    return round(candidate, 2)


def _ordered_actions(actions: list[tuple[Campaign, str]]) -> list[tuple[Campaign, str]]:
    # Order by the midnight when each action takes effect. An END for D and a
    # START for D+1 share a boundary, with END first so the next promo can layer
    # over the restored state.
    def event_key(item):
        campaign, action = item
        action = action.upper()
        try:
            event_date = date.fromisoformat(campaign.end_date if action == "END" else campaign.start_date)
            if action == "END":
                event_date += timedelta(days=1)
        except (TypeError, ValueError):
            event_date = date.max
        return (event_date, 0 if action == "END" else 1, int(campaign.id or 0))
    return sorted(actions, key=event_key)


def _promo_stack(raw) -> list[dict]:
    try:
        value = json.loads(_text(raw) or "[]")
        return [item for item in value if isinstance(item, dict)] if isinstance(value, list) else []
    except (TypeError, ValueError):
        return []


def _promo_frame(state: dict, owner: str = "", end_date: str = "") -> dict:
    return {
        "promo_code": state.get("promo_code", ""),
        "promo_type": state.get("promo_type", ""),
        "promo_value": state.get("promo_value", 0.0),
        "net_eur": state.get("net_eur"),
        "net_usd": state.get("net_usd"),
        "promo_label": state.get("promo_label", ""),
        "promo_owner": str(owner or ""),
        "end_date": str(end_date or ""),
        "list_eur": state.get("list_eur"),
        "list_usd": state.get("list_usd"),
        "original_list_eur": state.get("original_list_eur"),
        "original_list_usd": state.get("original_list_usd"),
        "price_owner": state.get("price_owner", ""),
        "allow_below_cost": bool(state.get("allow_below_cost", False)),
    }


def _infer_campaign_owner(campaigns, code: str, row, header: dict[str, int], when: str):
    """Recover the owner for rows written by older planner versions."""
    if not code or not campaigns:
        return None
    candidates = []
    for campaign in campaigns:
        if campaign.promo_code != code or not _matches(row, header, campaign):
            continue
        if (not campaign.start_date or campaign.start_date <= when) and campaign.end_date >= when:
            candidates.append(campaign)
    return max(candidates, key=lambda c: (c.start_date, int(c.id or 0))) if candidates else None


def _state_net(state, currency):
    base = state["list_eur" if currency == "EUR" else "list_usd"]
    if base is None:
        return None
    if not state.get("promo_code"):
        return round(float(base), 2)
    value = state.get(f"net_{currency.lower()}") if state.get("promo_type") == "net" else state.get("promo_value", 0)
    return round_promo_price(base, state.get("promo_type", ""), value if value is not None else base)


def _apply_monotonic_adjustments(states, rows, value, campaigns, boundary, floor_choices, floor_issues, review_rows=None):
    """Overlay PA after regular offers and floor decisions; never alter a base price."""
    enabled = {str(int(c.id or 0)): c for c in campaigns
               if c.monotonic_destinations and c.end_date >= boundary
               and (not c.start_date or c.start_date <= boundary)}
    if not enabled:
        return
    groups = {}
    for index, row in enumerate(rows):
        key = (_text(value(row, "ISO")).upper(), _text(value(row, "Plan")))
        groups.setdefault(key, {}).setdefault(int(_num(value(row, "Days")) or 0), []).append(index)
    for (iso, plan), durations in groups.items():
        if not any(iso in c.monotonic_destinations and plan in c.plans for c in enabled.values()):
            continue
        # Broad alignment is still owned by an active, opted-in campaign.
        # A manual offer can supply its net price without owning the PA lifetime.
        broad = [c for c in enabled.values() if c.monotonic_all_promos
                 and iso in c.monotonic_destinations and plan in c.plans]
        broad_owner = max(broad, key=lambda c: (c.start_date, int(c.id or 0))) if broad else None
        caps = {}
        for days in sorted(durations, reverse=True):
            next_caps = dict(caps)
            for index in durations[days]:
                state, row = states[index], rows[index]
                anchor = enabled.get(str(state.get("promo_owner", "")))
                if anchor and (iso not in anchor.monotonic_destinations or days not in anchor.days or plan not in anchor.plans):
                    anchor = None
                if anchor is None and broad_owner and state.get("promo_code"):
                    anchor = broad_owner
                nets = {currency: _state_net(state, currency) for currency in ("EUR", "USD")}
                targets = dict(nets)
                owners = {}
                for currency, current in nets.items():
                    if current is None:
                        continue
                    cap = caps.get(currency)
                    if cap and current > cap[0] + 1e-9:
                        targets[currency], owners[currency] = cap
                    # Unrelated cheaper points must not become new anchors in
                    # campaign-only mode. An adjusted point only carries its cap.
                    candidate_owner = owners.get(currency) or anchor
                    if candidate_owner:
                        candidate = (targets[currency], candidate_owner)
                        if currency not in next_caps or candidate[0] < next_caps[currency][0] - 1e-9:
                            next_caps[currency] = candidate
                if not owners:
                    continue
                # A separate underlying snapshot avoids confusing a synthetic PA
                # with the real campaign stack and supports independent currencies.
                state["pa_underlying"] = json.dumps({k: v for k, v in state.items() if k != "events"}, ensure_ascii=False)
                owner = next(iter(owners.values()))
                state.update(promo_code="PA", promo_type="net", promo_value=targets.get("EUR") or 0.,
                             promo_label="Promo adjustment", promo_owner=str(int(owner.id or 0)),
                             net_eur=targets["EUR"], net_usd=targets["USD"], pa_unresolved=False)
                state["events"].append(f"PA · {owner.name}")
                # Approvals belong to the source campaign and shorter duration.
                # Raising a PA row alone would break the promised monotonicity.
                below = False
                all_below_approved = True
                for currency, target in targets.items():
                    floor = _num(value(row, f"CostFloor_{currency}"))
                    if target is None or floor is None or target >= floor - 1e-9:
                        continue
                    below = True
                    source = owners.get(currency, owner)
                    choice = floor_choices.get((int(source.id or 0), iso, plan, days, currency))
                    if review_rows is None or (iso, plan, days) in review_rows:
                        floor_issues.append(CostFloorIssue(int(source.id or 0), source.name, iso, plan, days, currency,
                                                         state["list_eur" if currency == "EUR" else "list_usd"],
                                                         target, floor, floor, state["list_eur" if currency == "EUR" else "list_usd"], True))
                    if choice != "override":
                        all_below_approved = False
                        if review_rows is None or (iso, plan, days) in review_rows:
                            state["pa_unresolved"] = True
                state["allow_below_cost"] = below and all_below_approved
            caps = next_caps


def preview_batch(
    path: Path | pd.DataFrame,
    actions: list[tuple[Campaign, str]],
    execution_date: str = "",
    *,
    adjust_below_floor_to_margin: bool = False,
    cost_floor_choices: dict[tuple[int, str, str, int, str], str] | None = None,
    conflict_choices: dict[tuple[int, str, str, int], str] | None = None,
    date_price_changes: list[dict] | None = None,
    campaigns: list[Campaign] | None = None,
    row_index: dict[tuple[str, str, int], list[int]] | None = None,
    row_index_signature: tuple[int, int] | None = None,
    review_rows: set[tuple[str, str, int]] | None = None,
    adjustment_review_rows: set[tuple[str, str, int]] | None = None,
) -> BatchPreview:
    """Simulate a whole execution batch without writing the workbook.

    The simulation is sequential (END before START), so hand-overs on the same
    pricing row are represented as one final before/after price change.
    """
    ordered = _ordered_actions(actions)
    if campaigns is None:
        campaigns = list({int(c.id or 0): c for c, _ in ordered}.values())
    is_frame = isinstance(path, pd.DataFrame)
    workbook_path = None if is_frame else Path(path)
    wb = None
    if is_frame:
        frame = path.astype(object).where(pd.notna(path), None).reset_index(drop=True)
        h = {str(column): index + 1 for index, column in enumerate(frame.columns)}
        rows = list(frame.itertuples(index=False, name=None))
        promo_catalogue = {}
    else:
        # Workbook preview remains supported for legacy callers. The shared
        # timeline path supplies a DataFrame and avoids repeated XLSX reads.
        wb = load_workbook(workbook_path, read_only=True, data_only=False)
    try:
        if wb is not None:
            ws = wb[PRICING_SHEET] if PRICING_SHEET in wb.sheetnames else wb[wb.sheetnames[0]]
            h = _headers(ws)
            rows = list(ws.iter_rows(min_row=2, values_only=True))
            promo_catalogue = {p.code: p for p in _read_promos(wb)}
        for col in ("ISO", "Plan", "Days", "PromoCode"):
            if col not in h:
                raise ValueError(f"Pricing workbook is missing {col}.")

        # Historical decisions still project forward; issue reporting can be
        # restricted to points affected at the selected date.
        normalized_floor_choices = {
            (int(cid), str(iso).upper(), str(plan), int(days), str(currency).upper()): str(choice).lower()
            for (cid, iso, plan, days, currency), choice in (cost_floor_choices or {}).items()
        }
        normalized_conflict_choices = {
            (int(cid), str(iso).upper(), str(plan), int(days)): str(choice).lower()
            for (cid, iso, plan, days), choice in (conflict_choices or {}).items()
        }
        normalized_date_prices = {}
        start_price_keys = set()
        for change in date_price_changes or []:
            row_key = change.get("row", change)
            try:
                campaign_id = int(change.get("campaign_id") or 0)
                iso = str(row_key.get("iso", "")).upper()
                plan = str(row_key.get("plan", ""))
                days = int(row_key.get("days") or 0)
                unit = str(row_key.get("pricing_unit_id", ""))
            except (TypeError, ValueError):
                continue
            for currency in ("EUR", "USD"):
                value = _num(change.get(f"price_{currency.lower()}"))
                if value is not None:
                    normalized_date_prices[(campaign_id, iso, plan, days, unit, currency)] = float(value)
                    if change.get("apply_on_start"):
                        start_price_keys.add((campaign_id, iso, plan, days, unit, currency))
        collision_issues: list[PromoCollisionIssue] = []

        def floor_choice(campaign_id, iso, plan, days, currency):
            return normalized_floor_choices.get(
                (int(campaign_id or 0), str(iso).upper(), str(plan), int(days), str(currency).upper())
            )

        def row_value(values, name):
            return values[h[name] - 1] if name in h else None

        def promo_meta(values, code: str):
            code = _text(code)
            opt = promo_catalogue.get(code)
            if opt is not None:
                return opt.promo_type, float(opt.value), opt.label or code
            # Legacy/static fallback when an existing code is not in the hidden catalogue.
            ptype = _text(row_value(values, "PromoType")).lower()
            pvalue = _num(row_value(values, "PromoValue")) or 0.0
            plabel = _text(row_value(values, "PromoName")) or _text(row_value(values, "PromoLabel")) or code
            if ptype in {"percentage", "pct", "%"}:
                ptype = "percent"
            if ptype in {"fixed", "amount", "value"}:
                ptype = "absolute"
            return ptype, float(pvalue), plabel or code

        states = []
        for values in rows:
            code = _text(row_value(values, "PromoCode"))
            ptype, pvalue, plabel = promo_meta(values, code) if code else ("", 0.0, "")
            eur = _num(row_value(values, "Price_EUR"))
            usd = _num(row_value(values, "Price_USD"))
            states.append({
                "promo_code": code,
                "promo_type": ptype,
                "promo_value": pvalue,
                "promo_label": plabel,
                "list_eur": eur,
                "list_usd": usd,
                "original_list_eur": _num(row_value(values, ORIGINAL_PRICE_EUR_COLUMN)),
                "original_list_usd": _num(row_value(values, ORIGINAL_PRICE_USD_COLUMN)),
                "price_owner": _text(row_value(values, PRICE_OWNER_COLUMN)),
                "promo_owner": _text(row_value(values, PROMO_OWNER_COLUMN)),
                "promo_stack": _promo_stack(row_value(values, PROMO_STACK_COLUMN)),
                "allow_below_cost": str(row_value(values, "AllowBelowCost") or "").strip().lower() in {"true", "1", "yes"},
                "net_eur": _num(row_value(values, "PromoNet_EUR")),
                "net_usd": _num(row_value(values, "PromoNet_USD")),
                "manual_pa_links": _text(row_value(values, "_ManualPALinks")),
                "pa_underlying": _text(row_value(values, "_PromoPlannerPAUnderlying")),
                "pa_unresolved": str(row_value(values, "_PromoPlannerPAUnresolved") or "").lower() in {"true", "1"},
                "events": [],
            })

        original = [dict(x) for x in states]
        for state in states:
            if state["pa_underlying"]:
                try:
                    underlying = json.loads(state["pa_underlying"])
                except (TypeError, ValueError):
                    raise ValueError("A PA row has invalid restoration data; restore it from a pricing snapshot.")
                state.update(underlying)
                state.update(pa_underlying="", pa_unresolved=False)
                state["events"] = ["REMOVE PA"]
        # Permanent list-price updates also replace the base hidden underneath
        # active/suspended promotions. Ending a promotion must reveal this new
        # base, rather than restoring the price from before the update.
        for values, state in zip(rows, states):
            iso = _text(row_value(values, "ISO")).upper()
            plan = _text(row_value(values, "Plan"))
            day = int(_num(row_value(values, "Days")) or 0)
            unit = _text(row_value(values, "PricingUnitIdUsed"))
            for currency in ("EUR", "USD"):
                base = normalized_date_prices.get((0, iso, plan, day, unit, currency))
                if base is None:
                    continue
                list_key = "list_eur" if currency == "EUR" else "list_usd"
                original_key = "original_list_eur" if currency == "EUR" else "original_list_usd"
                state[list_key] = base
                if state.get(original_key) is not None:
                    state[original_key] = base
                for suspended in state["promo_stack"]:
                    suspended[list_key] = base
                    if suspended.get(original_key) is not None:
                        suspended[original_key] = base
        for state in states:
            if state["allow_below_cost"]:
                state["allow_below_cost"] = False
                state["events"].append("RESET below-floor approval")
        items: list[BatchPreviewItem] = []
        floor_issues: list[CostFloorIssue] = []
        campaign_rules = []
        for c, action in ordered:
            action = action.upper()
            matched = changed = below = 0
            collisions: list[str] = []
            rule = {
                "campaign": c,
                "action": action,
                "destinations": set(c.destinations),
                "plans": set(c.plans),
                "days": set(c.days),
                "changed": 0,
            }
            campaign_rules.append(rule)
            for idx, values in enumerate(rows):
                if not _matches(values, h, c):
                    continue
                matched += 1
                state = states[idx]
                current = _text(state["promo_code"])
                iso = _text(row_value(values, "ISO"))
                plan = _text(row_value(values, "Plan"))
                day = int(_num(row_value(values, "Days")) or 0)

                if action == "END":
                    stack = state["promo_stack"]
                    owner = state["promo_owner"]
                    owner_matches = owner == str(int(c.id or 0))
                    if owner_matches or (not owner and current == c.promo_code):
                        # If this planner temporarily raised list prices when the promo
                        # started, restore the exact pre-campaign prices at END.
                        if state.get("price_owner") == c.promo_code:
                            if state.get("original_list_eur") is not None:
                                state["list_eur"] = state["original_list_eur"]
                            if state.get("original_list_usd") is not None:
                                state["list_usd"] = state["original_list_usd"]
                            state["original_list_eur"] = None
                            state["original_list_usd"] = None
                            state["price_owner"] = ""
                        state["allow_below_cost"] = False
                        resumed = None
                        while stack:
                            frame = stack.pop()
                            frame_end = _text(frame.get("end_date"))
                            # A suspended promo that ended during the override is
                            # removed instead of being resumed.
                            if frame_end and frame_end <= execution_date:
                                continue
                            resumed = frame
                            break
                        if resumed:
                            state.update({
                                "promo_code": _text(resumed.get("promo_code")),
                                "promo_type": _text(resumed.get("promo_type")),
                                "promo_value": float(resumed.get("promo_value") or 0.0),
                                "promo_label": _text(resumed.get("promo_label")),
                                "promo_owner": _text(resumed.get("promo_owner")),
                                "list_eur": _num(resumed.get("list_eur")),
                                "list_usd": _num(resumed.get("list_usd")),
                                "original_list_eur": _num(resumed.get("original_list_eur")),
                                "original_list_usd": _num(resumed.get("original_list_usd")),
                                "price_owner": _text(resumed.get("price_owner")),
                                "allow_below_cost": bool(resumed.get("allow_below_cost", False)),
                            })
                        else:
                            state.update({
                                "promo_code": "", "promo_type": "", "promo_value": 0.0,
                                "promo_label": "", "promo_owner": "",
                            })
                        state["events"].append(f"END {c.name}")
                        changed += 1
                        rule["changed"] += 1
                    else:
                        # If this campaign is underneath another active promo,
                        # expire its suspended layer so it cannot return later.
                        kept = [frame for frame in stack if str(frame.get("promo_owner", "")) != str(int(c.id or 0))]
                        if len(kept) != len(stack):
                            state["promo_stack"] = kept
                            state["events"].append(f"EXPIRE {c.name}")

                elif action == "START":
                    unresolved_conflict = False
                    if current and current != c.promo_code:
                        collision_key = (int(c.id or 0), iso.upper(), plan, day)
                        resolution = normalized_conflict_choices.get(collision_key)
                        collision_issues.append(PromoCollisionIssue(
                            int(c.id or 0), c.name, iso, plan, day, current,
                            _text(state["promo_label"]) or current,
                        ))
                        if resolution == "keep":
                            state["events"].append(f"KEEP {state['promo_label'] or current} instead of {c.name}")
                            continue
                        if resolution != "overwrite":
                            collisions.append(f"{c.name}: {iso} / {plan} / {day}d already has {current}")
                            unresolved_conflict = True
                    if unresolved_conflict:
                        continue
                    if current and current != c.promo_code:
                        owner = state.get("promo_owner", "")
                        owned_campaign = next((x for x in (campaigns or []) if str(int(x.id or 0)) == owner), None)
                        if not owned_campaign and not owner:
                            owned_campaign = _infer_campaign_owner(campaigns, current, values, h, execution_date)
                        state["promo_stack"].append(_promo_frame(
                            state,
                            str(int(owned_campaign.id or 0)) if owned_campaign else owner,
                            owned_campaign.end_date if owned_campaign else "",
                        ))
                        state["original_list_eur"] = None
                        state["original_list_usd"] = None
                        state["price_owner"] = ""
                        state["allow_below_cost"] = False
                    if current != c.promo_code and not unresolved_conflict:
                        changed += 1
                        rule["changed"] += 1
                    if unresolved_conflict:
                        continue
                    state["promo_code"] = c.promo_code
                    state["promo_type"] = c.promo_type
                    state["promo_value"] = float(c.promo_value)
                    state["promo_label"] = c.promo_name or c.promo_code
                    state["promo_owner"] = str(int(c.id or 0))
                    state["allow_below_cost"] = False
                    state["events"].append(f"START {c.name}")
                    unit = _text(row_value(values, "PricingUnitIdUsed"))
                    for currency in ("EUR", "USD"):
                        price_key = (int(c.id or 0), iso.upper(), plan, day, unit, currency)
                        if price_key in start_price_keys:
                            list_key = "list_eur" if currency == "EUR" else "list_usd"
                            original_key = "original_list_eur" if currency == "EUR" else "original_list_usd"
                            if state.get(original_key) is None:
                                state[original_key] = state.get(list_key)
                            state[list_key] = normalized_date_prices[price_key]
                            state["price_owner"] = c.promo_code
                else:
                    raise ValueError(f"Unsupported batch action: {action}")
            items.append(BatchPreviewItem(int(c.id or 0), c.name, action, matched, changed, below, collisions[:100]))

        # Saved editor changes are dated list-price overrides. They may belong
        # to an active promo (and expire with it) or be independent price edits.
        for values, state in zip(rows, states):
            iso = _text(row_value(values, "ISO")).upper()
            plan = _text(row_value(values, "Plan"))
            day = int(_num(row_value(values, "Days")) or 0)
            unit = _text(row_value(values, "PricingUnitIdUsed"))
            try:
                owner = int(state.get("promo_owner") or 0)
            except (TypeError, ValueError):
                owner = 0
            for currency in ("EUR", "USD"):
                scheduled = normalized_date_prices.get(
                    (owner, iso, plan, day, unit, currency),
                    normalized_date_prices.get((0, iso, plan, day, unit, currency)),
                )
                if scheduled is None:
                    continue
                list_key = "list_eur" if currency == "EUR" else "list_usd"
                original_key = "original_list_eur" if currency == "EUR" else "original_list_usd"
                is_campaign_override = (owner, iso, plan, day, unit, currency) in normalized_date_prices and owner != 0
                if is_campaign_override and state.get(original_key) is None and state.get("promo_code"):
                    state[original_key] = state.get(list_key)
                if is_campaign_override and state.get("promo_code"):
                    state["price_owner"] = state["promo_code"]
                state[list_key] = scheduled
                state["events"].append(f"SCHEDULED PRICE EDIT ({currency})")

        # Manual/detailed promo overrides must participate in the same floor
        # and monotonicity calculation as regular campaigns, including PA nets.
        manual_changes = [change for change in (date_price_changes or [])
                          if any(field in change for field in ("promo_code", "promo_type", "promo_value", "promo_net_eur", "promo_net_usd", "manual_pa_links"))]
        manual_indexes = {}
        for idx, values in (enumerate(rows) if manual_changes else ()):

            manual_key = (_text(row_value(values, "ISO")).upper(), _text(row_value(values, "Plan")),
                          int(_num(row_value(values, "Days")) or 0), _text(row_value(values, "PricingUnitIdUsed")))
            manual_indexes.setdefault(manual_key, []).append(idx)
        manual_rows = set()
        active_owners = {str(int(c.id or 0)): c for c in (campaigns or [])}
        for change in manual_changes:
            identity = change.get("row", change)
            manual_key = (str(identity.get("iso", "")).upper(), str(identity.get("plan", "")),
                          int(identity.get("days") or 0), str(identity.get("pricing_unit_id", "")))
            for idx in manual_indexes.get(manual_key, []):
                state = states[idx]
                owner = active_owners.get(str(state.get("promo_owner", "")))
                if (change.get("is_base_price_change") and owner
                        and owner.start_date >= change.get("effective_from", "")):
                    continue  # A later timed campaign overlays the manual offer.
                manual_rows.add(idx)
                for source, target in (("promo_code", "promo_code"), ("promo_type", "promo_type"),
                                       ("promo_value", "promo_value"), ("promo_label", "promo_label"),
                                       ("promo_net_eur", "net_eur"), ("promo_net_usd", "net_usd"),
                                       ("manual_pa_links", "manual_pa_links")):
                    if source in change:
                        state[target] = change[source]
                state.update(pa_underlying="", pa_unresolved=False)
                if change.get("is_base_price_change"):
                    state["promo_owner"] = ""
                state["events"].append("MANUAL PROMO EDIT")

        # Review only affected points when the timeline supplies a date scope.
        # Existing saved campaign decisions still carry forward until expiry.
        floor_issues = []
        below_by_campaign: dict[int, set[tuple[str, str, int]]] = {}
        campaign_by_id = {str(int(c.id or 0)): c for c in (campaigns or [])}
        for idx, values in enumerate(rows):
            state = states[idx]
            state["allow_below_cost"] = False
            code = _text(state["promo_code"])
            iso = _text(row_value(values, "ISO")).upper()
            plan = _text(row_value(values, "Plan"))
            day = int(_num(row_value(values, "Days")) or 0)
            in_review = review_rows is None or (iso, plan, day) in review_rows
            if not code:
                continue
            active_campaign = campaign_by_id.get(str(state.get("promo_owner", "")))
            if active_campaign is None:
                resulting_date = (date.fromisoformat(execution_date) + timedelta(days=1)).isoformat() if execution_date else execution_date
                active_campaign = _infer_campaign_owner(campaigns, code, values, h, resulting_date)
                if active_campaign is not None:
                    state["promo_owner"] = str(int(active_campaign.id or 0))
            campaign_id = int(active_campaign.id or 0) if active_campaign else 0
            campaign_name = active_campaign.name if active_campaign else (state["promo_label"] or code or "Price change")
            pa_approvals = []
            for currency in ("EUR", "USD"):
                list_key = "list_eur" if currency == "EUR" else "list_usd"
                floor = _num(row_value(values, f"CostFloor_{currency}")) if f"CostFloor_{currency}" in h else None
                base = state[list_key]
                if base is None or floor is None:
                    continue
                net_price = _state_net(state, currency)
                if net_price >= floor - 1e-9:
                    continue
                target = float(floor) * 1.10
                required = float(base) if state["promo_type"] == "net" else required_list_price_for_target(target, state["promo_type"], state["promo_value"])
                if in_review:
                    floor_issues.append(CostFloorIssue(
                        campaign_id, campaign_name, iso, plan, day, currency,
                        float(base), float(net_price), float(floor), float(target), float(required),
                        code == "PA" and state["promo_type"] == "net"
                    ))
                below_by_campaign.setdefault(campaign_id, set()).add((iso, plan, day))
                choice = floor_choice(campaign_id, iso, plan, day, currency) if campaign_id or in_review or idx in manual_rows else None
                if code == "PA" and state["promo_type"] == "net":
                    pa_approvals.append(choice == "override")
                if choice == "raise" and not (code == "PA" and state["promo_type"] == "net"):
                    original_key = "original_list_eur" if currency == "EUR" else "original_list_usd"
                    if state.get(original_key) is None:
                        state[original_key] = float(base)
                    if not state.get("price_owner"):
                        state["price_owner"] = code
                    state[list_key] = float(required)
                    state["events"].append(f"COST FLOOR PRICE RAISED ({currency})")
                elif choice == "override":
                    if not state["allow_below_cost"]:
                        state["events"].append(f"BELOW-COST OVERRIDE ({currency})")
                    state["allow_below_cost"] = True
            if pa_approvals:
                state["allow_below_cost"] = all(pa_approvals)
                state["pa_unresolved"] = in_review and not all(pa_approvals)
        for item in items:
            item.below_floor_rows = len(below_by_campaign.get(item.campaign_id, set()))

        boundary = (date.fromisoformat(execution_date) + timedelta(days=1)).isoformat() if execution_date else date.today().isoformat()
        _apply_monotonic_adjustments(states, rows, row_value, campaigns or [], boundary,
                                    normalized_floor_choices, floor_issues,
                                    adjustment_review_rows if adjustment_review_rows is not None else review_rows)

        changes: list[PriceChange] = []
        for idx, values in enumerate(rows):
            before = original[idx]
            after = states[idx]
            if not after["events"]:
                continue

            iso = _text(row_value(values, "ISO"))
            country = _text(row_value(values, "Country")) or iso
            changes.append(PriceChange(
                iso=iso,
                country=country,
                plan=_text(row_value(values, "Plan")),
                days=int(_num(row_value(values, "Days")) or 0),
                change=" → ".join(after["events"]),
                promo_before=before["promo_label"] or before["promo_code"] or "—",
                promo_after=after["promo_label"] or after["promo_code"] or "—",
                list_eur_before=before["list_eur"],
                list_eur_after=after["list_eur"],
                net_eur_before=_state_net(before, "EUR"),
                net_eur_after=_state_net(after, "EUR"),
                list_usd_before=before["list_usd"],
                list_usd_after=after["list_usd"],
                net_usd_before=_state_net(before, "USD"),
                net_usd_after=_state_net(after, "USD"),
                allow_below_cost_before=bool(before["allow_below_cost"]),
                allow_below_cost_after=bool(after["allow_below_cost"]),
            ))

        projected_rows = []
        for values, state in zip(rows, states):
            projected = {
                name: values[index - 1] if index - 1 < len(values) else None
                for name, index in h.items()
            }
            projected.update({
                "PromoCode": state["promo_code"],
                "PromoType": state["promo_type"],
                "PromoValue": state["promo_value"],
                "PromoLabel": state["promo_label"],
                "PromoName": state["promo_label"],
                "PromoNet_EUR": state.get("net_eur") if state["promo_type"] == "net" else None,
                "PromoNet_USD": state.get("net_usd") if state["promo_type"] == "net" else None,
                "_ManualPALinks": state.get("manual_pa_links", ""),
                "_PromoPlannerPAUnderlying": state["pa_underlying"],
                "_PromoPlannerPAUnresolved": state["pa_unresolved"],
                "Price_EUR": state["list_eur"],
                "Price_USD": state["list_usd"],
                "AllowBelowCost": bool(state["allow_below_cost"]),
                ORIGINAL_PRICE_EUR_COLUMN: state["original_list_eur"],
                ORIGINAL_PRICE_USD_COLUMN: state["original_list_usd"],
                PRICE_OWNER_COLUMN: state["price_owner"],
                PROMO_OWNER_COLUMN: state["promo_owner"],
                PROMO_STACK_COLUMN: json.dumps(state["promo_stack"], ensure_ascii=False) if state["promo_stack"] else "",
            })
            projected_rows.append(projected)

        return BatchPreview(execution_date, items, floor_issues, collision_issues, changes, projected_rows)
    finally:
        if wb is not None:
            wb.close()


def execute_batch(
    path: Path,
    actions: list[tuple[Campaign, str]],
    backup_dir: Path,
    execution_date: str,
    *,
    adjust_below_floor_to_margin: bool = False,
    cost_floor_choices: dict[tuple[int, str, str, int, str], str] | None = None,
    conflict_choices: dict[tuple[int, str, str, int], str] | None = None,
    date_price_changes: list[dict] | None = None,
    campaigns: list[Campaign] | None = None,
    row_index: dict[tuple[str, str, int], list[int]] | None = None,
    row_index_signature: tuple[int, int] | None = None,
    progress_callback: Callable[[str, int, int], None] | None = None,
) -> tuple[Path, dict[tuple[int, str], int], int]:
    def report(message: str, value: int, maximum: int = 100) -> None:
        if progress_callback:
            progress_callback(message, value, maximum)

    report("Checking the final preview…", 3)
    ordered = _ordered_actions(actions)
    preview = preview_batch(
        path, ordered, execution_date,
        adjust_below_floor_to_margin=adjust_below_floor_to_margin,
        cost_floor_choices=cost_floor_choices,
        conflict_choices=conflict_choices,
        date_price_changes=date_price_changes,
        campaigns=campaigns,
        row_index=row_index,
        row_index_signature=row_index_signature,
    )
    if preview.collisions:
        raise RuntimeError("Promo collision detected. Resolve before execution:\n" + "\n".join(preview.collisions[:30]))
    normalized_floor_choices = {
        (int(cid), str(iso).upper(), str(plan), int(days), str(currency).upper()): str(choice).lower()
        for (cid, iso, plan, days, currency), choice in (cost_floor_choices or {}).items()
    }
    unresolved = [
        issue for issue in preview.cost_floor_issues
        if normalized_floor_choices.get((issue.campaign_id, issue.iso.upper(), issue.plan, issue.days, issue.currency.upper()))
        not in ({"override"} if issue.is_adjustment else {"raise", "override"})
    ]
    if unresolved:
        raise RuntimeError(f"Choose a cost-floor action for all {len(unresolved)} affected currency prices before applying.")
    normalized_conflict_choices = {
        (int(cid), str(iso).upper(), str(plan), int(days)): str(choice).lower()
        for (cid, iso, plan, days), choice in (conflict_choices or {}).items()
    }
    normalized_date_prices = {}
    for change in date_price_changes or []:
        row_key = change.get("row", change)
        try:
            campaign_id = int(change.get("campaign_id") or 0)
            iso = str(row_key.get("iso", "")).upper()
            plan = str(row_key.get("plan", ""))
            days = int(row_key.get("days") or 0)
            unit = str(row_key.get("pricing_unit_id", ""))
        except (TypeError, ValueError):
            continue
        for currency in ("EUR", "USD"):
            value = _num(change.get(f"price_{currency.lower()}"))
            if value is not None:
                normalized_date_prices[(campaign_id, iso, plan, days, unit, currency)] = float(value)
    if conflict_choices is not None:
        unresolved_collisions = [
            issue for issue in preview.collision_issues
            if normalized_conflict_choices.get((issue.campaign_id, issue.iso.upper(), issue.plan, issue.days))
            not in {"overwrite", "keep"}
        ]
        if unresolved_collisions:
            raise RuntimeError(f"Choose how to resolve all {len(unresolved_collisions)} promo conflicts before applying.")

    # Map exact affected pricing points/currencies to the required safe list price.
    adjustment_map: dict[tuple[int, str, str, int, str], float] = {}
    for issue in preview.cost_floor_issues:
        choice = normalized_floor_choices.get((issue.campaign_id, issue.iso.upper(), issue.plan, issue.days, issue.currency.upper()))
        if choice != "raise":
            continue
        key = (issue.campaign_id, issue.iso.upper(), issue.plan, issue.days, issue.currency)
        adjustment_map[key] = max(adjustment_map.get(key, 0.0), issue.required_list_price)

    # PA uses the same final rows as date projection. Preserve the workbook
    # container (including VBA) while writing exact computed values, rather
    # than reimplementing synthetic offers in the legacy per-action writer.
    if any(c.monotonic_destinations for c in (campaigns or [])) or any(
        row.get("PromoCode") == "PA" for row in preview.projected_rows
    ) or any("REMOVE PA" in item.change for item in preview.price_changes):
        try:
            from automation.pricing_book import normalize_pricing_dataframe
        except ModuleNotFoundError:
            from pricing_book import normalize_pricing_dataframe
        backup_dir.mkdir(parents=True, exist_ok=True)
        backup = backup_dir / f"{path.stem}_before_batch_{execution_date}_{datetime.now():%Y%m%d-%H%M%S-%f}{path.suffix}"
        shutil.copy2(path, backup)
        wb = load_workbook(path, keep_vba=path.suffix.lower() == ".xlsm")
        try:
            catalogue = [{"promo_code": p.code, "promo_type": p.promo_type, "promo_value": p.value, "label": p.label}
                         for p in _read_promos(wb)]
            catalogue += [{"promo_code": c.promo_code, "promo_type": c.promo_type, "promo_value": c.promo_value, "label": c.promo_name}
                          for c in (campaigns or [])]
            prepared = normalize_pricing_dataframe(pd.DataFrame(preview.projected_rows), promo_catalog=catalogue)
            ws = wb[PRICING_SHEET] if PRICING_SHEET in wb.sheetnames else wb[wb.sheetnames[0]]
            headers = _headers(ws)
            for column in prepared.columns:
                if column not in headers:
                    headers[column] = ws.max_column + 1
                    ws.cell(1, headers[column]).value = column
            for number, row in enumerate(prepared.to_dict(orient="records"), start=2):
                for column, value in row.items():
                    ws.cell(number, headers[column]).value = None if pd.isna(value) else value
            fd, temporary = tempfile.mkstemp(prefix=".pa-", suffix=path.suffix, dir=str(path.parent))
            os.close(fd)
            try:
                wb.save(temporary)
                os.replace(temporary, path)
            finally:
                if Path(temporary).exists():
                    Path(temporary).unlink()
        except Exception:
            shutil.copy2(backup, path)
            raise
        finally:
            wb.close()
        return backup, {(i.campaign_id, i.action): i.changed_rows for i in preview.items}, len(adjustment_map)

    stage_started = time.perf_counter()
    report("Creating a recovery copy…", 10)
    backup_dir.mkdir(parents=True, exist_ok=True)
    ts = datetime.now().strftime("%Y%m%d-%H%M%S-%f")
    backup = backup_dir / f"{path.stem}_before_batch_{execution_date}_{ts}{path.suffix}"
    shutil.copy2(path, backup)
    report(f"Recovery copy created in {time.perf_counter() - stage_started:.1f}s.", 12)

    report("Opening the workbook for safe updates…", 15)
    stage_started = time.perf_counter()
    wb = load_workbook(path, keep_vba=path.suffix.lower() == ".xlsm")
    report(f"Workbook opened in {time.perf_counter() - stage_started:.1f}s.", 18)
    changed_map: dict[tuple[int, str], int] = {}
    adjusted_cells = 0
    try:
        ws = wb[PRICING_SHEET] if PRICING_SHEET in wb.sheetnames else wb[wb.sheetnames[0]]
        h = _ensure_promo_name_column(ws, _headers(ws))
        h = _ensure_planner_price_columns(ws, h)
        if "AllowBelowCost" in h:
            for row_number in range(2, ws.max_row + 1):
                ws.cell(row_number, h["AllowBelowCost"]).value = False
        report("Preparing price changes…", 20)
        row_count = max(0, ws.max_row - 1)
        campaign_rules = []
        for c, action in ordered:
            action = action.upper()
            campaign_rules.append({
                "campaign": c,
                "action": action,
                "destinations": set(c.destinations),
                "plans": set(c.plans),
                "days": set(c.days),
                "changed": 0,
            })

        current_stat = path.stat()
        index_is_current = (
            row_index is not None
            and row_index_signature == (current_stat.st_mtime_ns, current_stat.st_size)
        )
        rules_by_row: dict[int, list[dict]] = {}
        if index_is_current:
            for rule in campaign_rules:
                for iso in rule["destinations"]:
                    for plan in rule["plans"]:
                        for day in rule["days"]:
                            for row_number in row_index.get((str(iso).upper(), str(plan), int(day)), ()):
                                rules_by_row.setdefault(row_number, []).append(rule)
            row_numbers = sorted(rules_by_row)
            report(f"Found {len(row_numbers):,} matching pricing row(s) from the saved index.", 25)
        else:
            # A stale or missing index is never trusted; fall back to the full scan.
            row_numbers = range(2, ws.max_row + 1)
            report("Locating matching rows in the workbook…", 20)

        row_scan_started = time.perf_counter()
        for row_position, r in enumerate(row_numbers, start=1):
            if not index_is_current:
                vals = [ws.cell(r, col).value for col in range(1, ws.max_column + 1)]
                iso = _text(vals[h["ISO"] - 1]).upper()
                plan = _text(vals[h["Plan"] - 1])
                day_value = _num(vals[h["Days"] - 1])
                if day_value is None:
                    continue
                day = int(day_value)
                row_rules = [
                    rule for rule in campaign_rules
                    if iso in rule["destinations"] and plan in rule["plans"] and day in rule["days"]
                ]
            else:
                row_rules = rules_by_row[r]
                iso = _text(ws.cell(r, h["ISO"]).value).upper()
                plan = _text(ws.cell(r, h["Plan"]).value)
                day = int(_num(ws.cell(r, h["Days"]).value) or 0)
                vals = [ws.cell(r, col).value for col in range(1, ws.max_column + 1)]
            if not row_rules:
                continue
            if row_position % 50 == 0 or row_position == len(row_numbers):
                report(
                    f"Applying matching rows: {row_position:,} of {len(row_numbers):,}…",
                    25 + int(50 * row_position / max(1, len(row_numbers))),
                )
            promo_cell = ws.cell(r, h["PromoCode"])
            name_cell = ws.cell(r, h[PROMO_NAME_COLUMN])
            current = _text(promo_cell.value)
            owner_cell = ws.cell(r, h[PROMO_OWNER_COLUMN])
            stack_cell = ws.cell(r, h[PROMO_STACK_COLUMN])
            promo_owner = _text(owner_cell.value)
            promo_stack = _promo_stack(stack_cell.value)
            for rule in row_rules:
                c = rule["campaign"]
                action = rule["action"]
                if action == "END":
                    is_top = promo_owner == str(int(c.id or 0)) or (not promo_owner and current == c.promo_code)
                    if is_top:
                        if _text(ws.cell(r, h[PRICE_OWNER_COLUMN]).value) == c.promo_code:
                            original_eur = _num(ws.cell(r, h[ORIGINAL_PRICE_EUR_COLUMN]).value)
                            original_usd = _num(ws.cell(r, h[ORIGINAL_PRICE_USD_COLUMN]).value)
                            if original_eur is not None and "Price_EUR" in h:
                                ws.cell(r, h["Price_EUR"]).value = original_eur
                            if original_usd is not None and "Price_USD" in h:
                                ws.cell(r, h["Price_USD"]).value = original_usd
                            ws.cell(r, h[ORIGINAL_PRICE_EUR_COLUMN]).value = None
                            ws.cell(r, h[ORIGINAL_PRICE_USD_COLUMN]).value = None
                            ws.cell(r, h[PRICE_OWNER_COLUMN]).value = None
                        if "AllowBelowCost" in h:
                            ws.cell(r, h["AllowBelowCost"]).value = False
                        resumed = None
                        while promo_stack:
                            frame = promo_stack.pop()
                            frame_end = _text(frame.get("end_date"))
                            if frame_end and frame_end <= execution_date:
                                continue
                            resumed = frame
                            break
                        if resumed:
                            promo_cell.value = _text(resumed.get("promo_code"))
                            name_cell.value = _text(resumed.get("promo_label"))
                            owner_cell.value = _text(resumed.get("promo_owner")) or None
                            promo_owner = _text(resumed.get("promo_owner"))
                            if "Price_EUR" in h:
                                ws.cell(r, h["Price_EUR"]).value = _num(resumed.get("list_eur"))
                            if "Price_USD" in h:
                                ws.cell(r, h["Price_USD"]).value = _num(resumed.get("list_usd"))
                            ws.cell(r, h[ORIGINAL_PRICE_EUR_COLUMN]).value = _num(resumed.get("original_list_eur"))
                            ws.cell(r, h[ORIGINAL_PRICE_USD_COLUMN]).value = _num(resumed.get("original_list_usd"))
                            ws.cell(r, h[PRICE_OWNER_COLUMN]).value = _text(resumed.get("price_owner")) or None
                            if "AllowBelowCost" in h:
                                ws.cell(r, h["AllowBelowCost"]).value = bool(resumed.get("allow_below_cost", False))
                        else:
                            promo_cell.value = ""
                            name_cell.value = ""
                            owner_cell.value = None
                            promo_owner = ""
                        current = _text(promo_cell.value)
                        rule["changed"] += 1
                    else:
                        kept = [frame for frame in promo_stack if str(frame.get("promo_owner", "")) != str(int(c.id or 0))]
                        if len(kept) != len(promo_stack):
                            promo_stack = kept
                            rule["changed"] += 1
                elif action == "START":
                    if current and current != c.promo_code:
                        resolution = normalized_conflict_choices.get((int(c.id or 0), iso, plan, day))
                        if resolution == "keep":
                            continue
                        if resolution != "overwrite":
                            raise RuntimeError(f"Collision while applying {c.name}: pricing row already has {current}")

                        owned_campaign = next(
                            (x for x in (campaigns or []) if str(int(x.id or 0)) == promo_owner), None
                        )
                        if not owned_campaign and not promo_owner:
                            owned_campaign = _infer_campaign_owner(campaigns, current, vals, h, execution_date)
                        existing_meta = next((p for p in _read_promos(wb) if p.code == current), None)
                        current_allow = bool(ws.cell(r, h["AllowBelowCost"]).value) if "AllowBelowCost" in h else False
                        frame_state = {
                            "promo_code": current,
                            "promo_type": existing_meta.promo_type if existing_meta else "",
                            "promo_value": existing_meta.value if existing_meta else 0.0,
                            "promo_label": _text(name_cell.value) or current,
                            "list_eur": _num(ws.cell(r, h["Price_EUR"]).value) if "Price_EUR" in h else None,
                            "list_usd": _num(ws.cell(r, h["Price_USD"]).value) if "Price_USD" in h else None,
                            "original_list_eur": _num(ws.cell(r, h[ORIGINAL_PRICE_EUR_COLUMN]).value),
                            "original_list_usd": _num(ws.cell(r, h[ORIGINAL_PRICE_USD_COLUMN]).value),
                            "price_owner": _text(ws.cell(r, h[PRICE_OWNER_COLUMN]).value),
                            "allow_below_cost": current_allow,
                        }
                        promo_stack.append(_promo_frame(
                            frame_state,
                            str(int(owned_campaign.id or 0)) if owned_campaign else promo_owner,
                            owned_campaign.end_date if owned_campaign else "",
                        ))
                        # Each active layer owns its own temporary price uplift.
                        ws.cell(r, h[ORIGINAL_PRICE_EUR_COLUMN]).value = None
                        ws.cell(r, h[ORIGINAL_PRICE_USD_COLUMN]).value = None
                        ws.cell(r, h[PRICE_OWNER_COLUMN]).value = None
                        if "AllowBelowCost" in h:
                            ws.cell(r, h["AllowBelowCost"]).value = False

                    row_adjusted = False
                    if adjust_below_floor_to_margin or cost_floor_choices is not None:
                        for currency in ("EUR", "USD"):
                            key = (int(c.id or 0), iso, plan, day, currency)
                            required = adjustment_map.get(key)
                            price_col = f"Price_{currency}"
                            if required is not None and price_col in h:
                                cell = ws.cell(r, h[price_col])
                                current_price = _num(cell.value)
                                if current_price is None or current_price < required - 1e-9:
                                    original_col = ORIGINAL_PRICE_EUR_COLUMN if currency == "EUR" else ORIGINAL_PRICE_USD_COLUMN
                                    original_cell = ws.cell(r, h[original_col])
                                    if _num(original_cell.value) is None and current_price is not None:
                                        original_cell.value = current_price
                                    ws.cell(r, h[PRICE_OWNER_COLUMN]).value = c.promo_code
                                    cell.value = required
                                    adjusted_cells += 1
                                row_adjusted = True

                    if current != c.promo_code:
                        promo_cell.value = c.promo_code
                        name_cell.value = c.promo_name
                        rule["changed"] += 1
                        current = c.promo_code
                    elif _text(name_cell.value) != c.promo_name:
                        name_cell.value = c.promo_name
                    promo_owner = str(int(c.id or 0))
                    owner_cell.value = str(int(c.id or 0))

                    if "AllowBelowCost" in h:
                        if cost_floor_choices is not None:
                            row_override = any(
                                normalized_floor_choices.get((int(c.id or 0), iso, plan, day, currency)) == "override"
                                for currency in ("EUR", "USD")
                            )
                            ws.cell(r, h["AllowBelowCost"]).value = row_override
                        else:
                            ws.cell(r, h["AllowBelowCost"]).value = False
            owner_cell.value = promo_owner or None
            stack_cell.value = json.dumps(promo_stack, ensure_ascii=False) if promo_stack else None
        if normalized_date_prices:
            for r in range(2, ws.max_row + 1):
                try:
                    iso = _text(ws.cell(r, h["ISO"]).value).upper()
                    plan = _text(ws.cell(r, h["Plan"]).value)
                    day = int(_num(ws.cell(r, h["Days"]).value) or 0)
                    unit = _text(ws.cell(r, h["PricingUnitIdUsed"]).value) if "PricingUnitIdUsed" in h else ""
                    owner = int(_text(ws.cell(r, h[PROMO_OWNER_COLUMN]).value) or 0)
                except (KeyError, TypeError, ValueError):
                    continue
                code = _text(ws.cell(r, h["PromoCode"]).value)
                for currency in ("EUR", "USD"):
                    scheduled = normalized_date_prices.get((owner, iso, plan, day, unit, currency))
                    price_col = f"Price_{currency}"
                    if scheduled is None or price_col not in h:
                        continue
                    if code:
                        original_col = ORIGINAL_PRICE_EUR_COLUMN if currency == "EUR" else ORIGINAL_PRICE_USD_COLUMN
                        original_cell = ws.cell(r, h[original_col])
                        current_price = _num(ws.cell(r, h[price_col]).value)
                        if _num(original_cell.value) is None and current_price is not None:
                            original_cell.value = current_price
                        ws.cell(r, h[PRICE_OWNER_COLUMN]).value = code
                    ws.cell(r, h[price_col]).value = scheduled
                    adjusted_cells += 1
        if "AllowBelowCost" in h:
            for row_number in range(2, ws.max_row + 1):
                ws.cell(row_number, h["AllowBelowCost"]).value = False
        promo_catalogue = {p.code: p for p in _read_promos(wb)}
        campaign_by_id = {str(int(c.id or 0)): c for c in (campaigns or [])}
        resulting_date = (date.fromisoformat(execution_date) + timedelta(days=1)).isoformat() if execution_date else execution_date
        for r in range(2, ws.max_row + 1):
            code = _text(ws.cell(r, h["PromoCode"]).value)
            if not code:
                continue
            iso = _text(ws.cell(r, h["ISO"]).value).upper()
            plan = _text(ws.cell(r, h["Plan"]).value)
            day = int(_num(ws.cell(r, h["Days"]).value) or 0)
            owner_id = _text(ws.cell(r, h[PROMO_OWNER_COLUMN]).value)
            active_campaign = campaign_by_id.get(owner_id)
            if active_campaign is None:
                row_values = [ws.cell(r, col).value for col in range(1, ws.max_column + 1)]
                active_campaign = _infer_campaign_owner(campaigns, code, row_values, h, resulting_date)
            campaign_id = int(active_campaign.id or 0) if active_campaign else 0
            promo = promo_catalogue.get(code)
            if promo is not None:
                promo_type = promo.promo_type
                promo_value = promo.value
            else:
                # Match preview's legacy fallback for older active promos whose
                # definition is stored on the pricing row rather than in the
                # hidden promo catalogue.
                promo_type = _text(ws.cell(r, h["PromoType"]).value).lower() if "PromoType" in h else ""
                promo_value = _num(ws.cell(r, h["PromoValue"]).value) or 0.0 if "PromoValue" in h else 0.0
                if promo_type in {"percentage", "pct", "%"}:
                    promo_type = "percent"
                if promo_type in {"fixed", "amount", "value"}:
                    promo_type = "absolute"
            any_override = False
            for currency in ("EUR", "USD"):
                price_col = f"Price_{currency}"
                floor_col = f"CostFloor_{currency}"
                if price_col not in h or floor_col not in h:
                    continue
                base = _num(ws.cell(r, h[price_col]).value)
                floor = _num(ws.cell(r, h[floor_col]).value)
                if base is None or floor is None:
                    continue
                if round_promo_price(base, promo_type, promo_value) >= floor - 1e-9:
                    continue
                choice = normalized_floor_choices.get((campaign_id, iso, plan, day, currency))
                if choice == "override":
                    any_override = True
                elif choice == "raise":
                    required = required_list_price_for_target(float(floor) * 1.10, promo_type, promo_value)
                    original_col = ORIGINAL_PRICE_EUR_COLUMN if currency == "EUR" else ORIGINAL_PRICE_USD_COLUMN
                    original_cell = ws.cell(r, h[original_col])
                    if _num(original_cell.value) is None:
                        original_cell.value = base
                    ws.cell(r, h[PRICE_OWNER_COLUMN]).value = code
                    ws.cell(r, h[price_col]).value = required
                    adjusted_cells += 1
                else:
                    raise RuntimeError(f"Missing cost-floor decision for {iso} / {plan} / {day}d / {currency}.")
            if "AllowBelowCost" in h:
                ws.cell(r, h["AllowBelowCost"]).value = any_override
        row_scan_duration = time.perf_counter() - row_scan_started
        for rule in campaign_rules:
            changed_map[(int(rule["campaign"].id or 0), rule["action"])] = rule["changed"]
        report(f"All date changes prepared in {row_scan_duration:.1f}s.", 75)
        try:
            wb.calculation.fullCalcOnLoad = True
            wb.calculation.forceFullCalc = True
            wb.calculation.calcMode = "auto"
        except Exception:
            pass
        report("Writing the complete Excel workbook…", 0, 0)
        save_started = time.perf_counter()
        fd, temp_name = tempfile.mkstemp(
            prefix=f".{path.stem}.planner-", suffix=path.suffix, dir=str(path.parent)
        )
        os.close(fd)
        temp_path = Path(temp_name)
        try:
            wb.save(temp_path)
            if not temp_path.exists() or temp_path.stat().st_size == 0:
                raise RuntimeError("The temporary price workbook was not written correctly.")
            report(f"Workbook written in {time.perf_counter() - save_started:.1f}s; verifying…", 0, 0)
            verify_started = time.perf_counter()
            check = load_workbook(temp_path, read_only=True, data_only=False)
            try:
                if not check.sheetnames:
                    raise RuntimeError("The temporary price workbook could not be verified.")
            finally:
                check.close()
            report(f"Workbook verified in {time.perf_counter() - verify_started:.1f}s.", 96)
            os.replace(temp_path, path)
        finally:
            if temp_path.exists():
                temp_path.unlink()
    except Exception:
        try:
            wb.close()
        except Exception:
            pass
        shutil.copy2(backup, path)
        raise
    finally:
        try:
            wb.close()
        except Exception:
            pass
    report("Price workbook saved successfully.", 100)
    return backup, changed_map, adjusted_cells


def restore_workbook_backup(path: Path, backup: Path) -> None:
    if not backup.exists():
        raise FileNotFoundError(f"Backup not found: {backup}")
    shutil.copy2(backup, path)
