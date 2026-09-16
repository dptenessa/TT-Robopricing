from __future__ import annotations

from pathlib import Path
from typing import Any, Iterable
import math

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.workbook.defined_name import DefinedName


PRICING_SHEET = "Pricing"
PROMOS_SHEET = "Promos"

# User-facing fields come first. Technical fields are retained but hidden.
PRICING_COLUMNS = [
    "Country",
    "ISO",
    "PricingUnitIdUsed",
    "Plan",
    "Days",
    "GB",
    "Price_EUR",
    "Price_USD",
    "PromoCode",
    "FinalPriceAfterPromo_EUR",
    "FinalPriceAfterPromo_USD",
    "CostFloor_EUR",
    "CostFloor_USD",
    "EUR_IsBelowCostFloor",
    "USD_IsBelowCostFloor",
    "AllowBelowCost",
    # Hidden technical fields below.
    "Provider",
    "ReferenceProvider",
    "ISO3",
    "PricingSourceUsed",
    "PricingRegionUsed",
    "PricingUnitCountriesUsed",
    "PromoScopeKey",
    "PromoType",
    "PromoValue",
    "PromoLabel",
    "PromoBasePrice",
]

HIDDEN_COLUMNS = {
    "Provider",
    "ReferenceProvider",
    "ISO3",
    "PricingSourceUsed",
    "PricingRegionUsed",
    "PricingUnitCountriesUsed",
    "PromoScopeKey",
    "PromoType",
    "PromoValue",
    "PromoLabel",
    "PromoBasePrice",
}

INPUT_COLUMNS = {"Price_EUR", "Price_USD", "PromoCode", "AllowBelowCost"}
FORMULA_COLUMNS = {
    "PromoType",
    "PromoValue",
    "PromoLabel",
    "PromoBasePrice",
    "FinalPriceAfterPromo_EUR",
    "FinalPriceAfterPromo_USD",
    "EUR_IsBelowCostFloor",
    "USD_IsBelowCostFloor",
}



def _text(value: Any) -> str:
    if value is None:
        return ""
    text = str(value).strip()
    return "" if text.lower() in {"nan", "none", "<na>", "nat"} else text

def _bool_value(value: Any) -> bool:
    if isinstance(value, bool):
        return value
    text = str(value if value is not None else "").strip().lower()
    return text in {"true", "t", "yes", "y", "1"}


def _num(value: Any) -> float | None:
    if value is None:
        return None
    if isinstance(value, str) and value.startswith("="):
        return None
    number = pd.to_numeric(value, errors="coerce")
    if pd.isna(number):
        return None
    return float(number)


def _promo_map_from_catalog(promo_catalog: Iterable[dict[str, Any]] | None) -> dict[str, dict[str, Any]]:
    out: dict[str, dict[str, Any]] = {}
    for promo in promo_catalog or []:
        code = _text(promo.get("promo_code", ""))
        if not code:
            continue
        promo_type = _text(promo.get("promo_type", "")).lower()
        if promo_type in {"percentage", "%"}:
            promo_type = "percent"
        out[code] = {
            "promo_code": code,
            "promo_type": promo_type,
            "promo_value": float(promo.get("promo_value", 0) or 0),
            "label": _text(promo.get("label", "")) or code,
        }
    return out


def _promo_map_from_workbook(wb) -> dict[str, dict[str, Any]]:
    if PROMOS_SHEET not in wb.sheetnames:
        return {}
    ws = wb[PROMOS_SHEET]
    out: dict[str, dict[str, Any]] = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row:
            continue
        code = _text(row[0] if len(row) > 0 else "")
        if not code:
            continue
        promo_type = _text(row[1] if len(row) > 1 else "").lower()
        if promo_type in {"percentage", "%"}:
            promo_type = "percent"
        value = _num(row[2] if len(row) > 2 else None) or 0.0
        label = _text(row[3] if len(row) > 3 else "") or code
        out[code] = {
            "promo_code": code,
            "promo_type": promo_type,
            "promo_value": float(value),
            "label": label,
        }
    return out


def _promo_net_price(price: Any, promo_type: str, promo_value: Any) -> float | None:
    base = _num(price)
    if base is None:
        return None
    ptype = str(promo_type or "").strip().lower()
    value = _num(promo_value) or 0.0
    if ptype == "percent":
        raw = base * (1.0 - value / 100.0)
    elif ptype == "absolute":
        raw = base - value
    else:
        return round(base * 20.0) / 20.0
    # Match EditorState.round_promo_price(): floor to 0.05.
    return math.floor(max(raw, 0.0) * 20.0 + 1e-12) / 20.0


def normalize_pricing_dataframe(
    df: pd.DataFrame,
    *,
    promo_catalog: Iterable[dict[str, Any]] | None = None,
    infer_legacy_allow: bool = True,
) -> pd.DataFrame:
    """Recompute derived pricing fields from list prices, promo code and floors.

    This function deliberately does not trust cached Excel formula values. It is
    used by partner export and workbook loading so manual Excel edits are safe
    even if Excel has not recalculated the file yet.
    """
    out = df.copy()
    out.columns = out.columns.astype(str).str.strip()

    for column in PRICING_COLUMNS:
        if column not in out.columns:
            out[column] = ""

    for column in (
        "Country", "ISO", "PricingUnitIdUsed", "Plan", "Provider", "ReferenceProvider",
        "ISO3", "PricingSourceUsed", "PricingRegionUsed", "PricingUnitCountriesUsed",
        "PromoScopeKey", "PromoCode", "PromoType", "PromoLabel",
    ):
        if column in out.columns:
            out[column] = out[column].map(_text)

    promo_map = _promo_map_from_catalog(promo_catalog)

    # The canonical workbook contains only T-Travel pricing rows. Provider is
    # retained for compatibility/audit, but blank legacy values must not cause
    # downstream consumers to discard a valid row.
    if "Provider" in out.columns:
        out["Provider"] = out["Provider"].map(_text)
        out.loc[out["Provider"].eq(""), "Provider"] = "HT"

    for currency in ("EUR", "USD"):
        price_col = f"Price_{currency}"
        floor_col = f"CostFloor_{currency}"
        out[price_col] = pd.to_numeric(out[price_col].where(~out[price_col].astype(str).str.startswith("="), pd.NA), errors="coerce")
        out[floor_col] = pd.to_numeric(out[floor_col].where(~out[floor_col].astype(str).str.startswith("="), pd.NA), errors="coerce")

    previous_blocked = out.get("IsPartnerExportBlocked", pd.Series(False, index=out.index)).map(_bool_value)
    previous_below_eur = out.get("EUR_IsBelowCostFloor", pd.Series(False, index=out.index)).map(_bool_value)
    previous_below_usd = out.get("USD_IsBelowCostFloor", pd.Series(False, index=out.index)).map(_bool_value)

    promo_types: list[str] = []
    promo_values: list[Any] = []
    promo_labels: list[str] = []
    final_eur: list[Any] = []
    final_usd: list[Any] = []

    for _, row in out.iterrows():
        code = _text(row.get("PromoCode", ""))
        promo = promo_map.get(code)
        if code and promo is None:
            # Legacy/static fallback if the code is not available in the hidden catalogue.
            ptype = _text(row.get("PromoType", "")).lower()
            pvalue = _num(row.get("PromoValue")) or 0.0
            plabel = _text(row.get("PromoLabel", "")) or code
            promo = {"promo_type": ptype, "promo_value": pvalue, "label": plabel}
        if not code:
            promo = None

        if promo:
            ptype = _text(promo.get("promo_type", "")).lower()
            pvalue = float(promo.get("promo_value", 0) or 0)
            plabel = _text(promo.get("label", "")) or code
        else:
            ptype, pvalue, plabel = "", "", ""

        promo_types.append(ptype)
        promo_values.append(pvalue)
        promo_labels.append(plabel)
        final_eur.append(_promo_net_price(row.get("Price_EUR"), ptype, pvalue))
        final_usd.append(_promo_net_price(row.get("Price_USD"), ptype, pvalue))

    out["PromoType"] = promo_types
    out["PromoValue"] = promo_values
    out["PromoLabel"] = promo_labels
    out["PromoBasePrice"] = [
        row_price if _text(code) else ""
        for row_price, code in zip(out["Price_EUR"], out["PromoCode"])
    ]
    out["FinalPriceAfterPromo_EUR"] = final_eur
    out["FinalPriceAfterPromo_USD"] = final_usd

    below_eur = (
        pd.to_numeric(out["FinalPriceAfterPromo_EUR"], errors="coerce").notna()
        & pd.to_numeric(out["CostFloor_EUR"], errors="coerce").notna()
        & (pd.to_numeric(out["FinalPriceAfterPromo_EUR"], errors="coerce") < pd.to_numeric(out["CostFloor_EUR"], errors="coerce") - 1e-9)
    )
    below_usd = (
        pd.to_numeric(out["FinalPriceAfterPromo_USD"], errors="coerce").notna()
        & pd.to_numeric(out["CostFloor_USD"], errors="coerce").notna()
        & (pd.to_numeric(out["FinalPriceAfterPromo_USD"], errors="coerce") < pd.to_numeric(out["CostFloor_USD"], errors="coerce") - 1e-9)
    )

    allow_col_present = "AllowBelowCost" in df.columns
    if allow_col_present:
        allow = out["AllowBelowCost"].map(_bool_value)
    elif infer_legacy_allow:
        # Preserve any legacy row that was below a floor but explicitly not blocked.
        allow = (previous_below_eur | previous_below_usd) & ~previous_blocked
    else:
        allow = pd.Series(False, index=out.index)

    blocked = (below_eur | below_usd) & ~allow
    reasons = []
    for e, u, b in zip(below_eur, below_usd, blocked):
        if not b:
            reasons.append("")
        elif e and u:
            reasons.append("EUR,USD")
        elif e:
            reasons.append("EUR")
        else:
            reasons.append("USD")

    out["EUR_IsBelowCostFloor"] = below_eur.astype(bool)
    out["USD_IsBelowCostFloor"] = below_usd.astype(bool)
    out["AllowBelowCost"] = allow.astype(bool)
    # Blocking is intentionally not persisted. Partner export derives it from
    # factual below-floor flags plus AllowBelowCost.
    out = out.drop(columns=["IsPartnerExportBlocked", "PartnerExportBlockReason"], errors="ignore")
    return out


def read_pricing_workbook(path: str | Path) -> pd.DataFrame:
    path = Path(path)
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet_name = PRICING_SHEET if PRICING_SHEET in wb.sheetnames else ("All_Data" if "All_Data" in wb.sheetnames else wb.sheetnames[0])
        ws = wb[sheet_name]
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            return pd.DataFrame()
        headers = [str(v if v is not None else "").strip() for v in rows[0]]
        data = [list(row[: len(headers)]) for row in rows[1:] if any(v is not None and str(v).strip() for v in row)]
        df = pd.DataFrame(data, columns=headers)
        # Recommendation is a disposable workbench annotation written after the
        # canonical workbook is generated. Never feed it back into pricing state.
        df = df.drop(columns=["Recommendation"], errors="ignore")
        promo_map = _promo_map_from_workbook(wb)
        return normalize_pricing_dataframe(df, promo_catalog=promo_map.values())
    finally:
        wb.close()


def _formula_refs(column_index: dict[str, int], row: int) -> dict[str, str]:
    return {name: f"{get_column_letter(idx)}{row}" for name, idx in column_index.items()}




def write_pricing_workbook(
    df: pd.DataFrame,
    path: str | Path,
    *,
    promo_catalog: Iterable[dict[str, Any]] | None = None,
) -> Path:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)

    promo_map = _promo_map_from_catalog(promo_catalog)
    normalized = normalize_pricing_dataframe(df, promo_catalog=promo_map.values())

    wb = Workbook()
    ws = wb.active
    ws.title = PRICING_SHEET
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(PRICING_COLUMNS))}{len(normalized) + 1}"

    # Excel should recalculate workbook formulas whenever it is opened.
    try:
        wb.calculation.fullCalcOnLoad = True
        wb.calculation.forceFullCalc = True
        wb.calculation.calcMode = "auto"
    except Exception:
        pass

    header_fill = PatternFill("solid", fgColor="E20074")
    header_font = Font(color="FFFFFF", bold=True)
    input_font = Font(color="0000FF")
    formula_font = Font(color="000000")
    linked_font = Font(color="008000")
    static_font = Font(color="666666")
    warning_fill = PatternFill("solid", fgColor="FCE4D6")
    blocked_fill = PatternFill("solid", fgColor="FFC7CE")
    allowed_fill = PatternFill("solid", fgColor="FFF2CC")

    for col_idx, name in enumerate(PRICING_COLUMNS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=name)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    column_index = {name: idx for idx, name in enumerate(PRICING_COLUMNS, start=1)}

    # Hidden promo catalogue used by dropdown + formulas.
    promo_ws = wb.create_sheet(PROMOS_SHEET)
    promo_ws.append(["PromoCode", "PromoType", "PromoValue", "PromoLabel"])
    for cell in promo_ws[1]:
        cell.fill = header_fill
        cell.font = header_font
    for promo in promo_map.values():
        promo_ws.append([
            promo["promo_code"],
            promo["promo_type"],
            promo["promo_value"],
            promo["label"],
        ])
    promo_ws.sheet_state = "hidden"

    if promo_map:
        promo_range = f"'{PROMOS_SHEET}'!$A$2:$A${len(promo_map) + 1}"
        try:
            wb.defined_names.add(DefinedName("PromoCodes", attr_text=promo_range))
        except AttributeError:
            wb.defined_names.append(DefinedName("PromoCodes", attr_text=promo_range))

    for excel_row, (_, row) in enumerate(normalized.iterrows(), start=2):
        refs = _formula_refs(column_index, excel_row)
        for name in PRICING_COLUMNS:
            col = column_index[name]
            cell = ws.cell(row=excel_row, column=col)

            if name == "PromoType":
                cell.value = f'=IF({refs["PromoCode"]}="","",IFERROR(VLOOKUP({refs["PromoCode"]},Promos!$A$2:$D$100,2,FALSE),""))'
            elif name == "PromoValue":
                cell.value = f'=IF({refs["PromoCode"]}="","",IFERROR(VLOOKUP({refs["PromoCode"]},Promos!$A$2:$D$100,3,FALSE),""))'
            elif name == "PromoLabel":
                cell.value = f'=IF({refs["PromoCode"]}="","",IFERROR(VLOOKUP({refs["PromoCode"]},Promos!$A$2:$D$100,4,FALSE),""))'
            elif name == "PromoBasePrice":
                cell.value = f'=IF({refs["PromoCode"]}="","",{refs["Price_EUR"]})'
            elif name in {"FinalPriceAfterPromo_EUR", "FinalPriceAfterPromo_USD"}:
                currency = name.rsplit("_", 1)[-1]
                price_ref = refs[f"Price_{currency}"]
                cell.value = (
                    f'=IF({refs["PromoCode"]}="",{price_ref},'
                    f'ROUNDDOWN(MAX(0,IF({refs["PromoType"]}="percent",'
                    f'{price_ref}*(1-{refs["PromoValue"]}/100),'
                    f'{price_ref}-{refs["PromoValue"]}))*20,0)/20)'
                )
            elif name == "EUR_IsBelowCostFloor":
                cell.value = f'=AND(ISNUMBER({refs["FinalPriceAfterPromo_EUR"]}),ISNUMBER({refs["CostFloor_EUR"]}),{refs["FinalPriceAfterPromo_EUR"]}<{refs["CostFloor_EUR"]}-0.000000001)'
            elif name == "USD_IsBelowCostFloor":
                cell.value = f'=AND(ISNUMBER({refs["FinalPriceAfterPromo_USD"]}),ISNUMBER({refs["CostFloor_USD"]}),{refs["FinalPriceAfterPromo_USD"]}<{refs["CostFloor_USD"]}-0.000000001)'
            else:
                value = row.get(name, "")
                if pd.isna(value):
                    value = ""
                if name == "AllowBelowCost":
                    value = bool(_bool_value(value))
                cell.value = value

            if name in INPUT_COLUMNS:
                cell.font = input_font
            elif name in FORMULA_COLUMNS:
                cell.font = formula_font
            elif name in {"Country", "ISO", "PricingUnitIdUsed", "Plan", "Days", "GB", "CostFloor_EUR", "CostFloor_USD"}:
                cell.font = linked_font
            else:
                cell.font = static_font

            if name in {"Price_EUR", "Price_USD", "FinalPriceAfterPromo_EUR", "FinalPriceAfterPromo_USD", "CostFloor_EUR", "CostFloor_USD"}:
                cell.number_format = '0.00'
            elif name in {"Days", "GB", "PromoValue"}:
                cell.number_format = '0.##'

    max_row = len(normalized) + 1
    if max_row >= 2:
        # Dropdown driven by hidden promo catalogue.
        if promo_map:
            dv_promo = DataValidation(type="list", formula1="=PromoCodes", allow_blank=True)
            dv_promo.error = "Choose a promo code from the list."
            dv_promo.errorTitle = "Invalid promo"
            dv_promo.prompt = "Select an approved promo or leave blank."
            dv_promo.promptTitle = "Promo"
            ws.add_data_validation(dv_promo)
            promo_letter = get_column_letter(column_index["PromoCode"])
            dv_promo.add(f"{promo_letter}2:{promo_letter}{max_row}")

        dv_allow = DataValidation(type="list", formula1='"TRUE,FALSE"', allow_blank=False)
        dv_allow.error = "Use TRUE or FALSE."
        dv_allow.errorTitle = "Invalid override"
        ws.add_data_validation(dv_allow)
        allow_letter = get_column_letter(column_index["AllowBelowCost"])
        dv_allow.add(f"{allow_letter}2:{allow_letter}{max_row}")

        # Conditional formatting: blocked = red, below-floor-but-overridden = yellow.
        eur_below_letter = get_column_letter(column_index["EUR_IsBelowCostFloor"])
        usd_below_letter = get_column_letter(column_index["USD_IsBelowCostFloor"])
        allow_letter = get_column_letter(column_index["AllowBelowCost"])
        visible_end = get_column_letter(column_index["AllowBelowCost"])
        target_range = f"A2:{visible_end}{max_row}"
        ws.conditional_formatting.add(
            target_range,
            FormulaRule(
                formula=[f'AND(OR(${eur_below_letter}2=TRUE,${usd_below_letter}2=TRUE),${allow_letter}2=FALSE)'],
                fill=blocked_fill,
            ),
        )
        ws.conditional_formatting.add(
            target_range,
            FormulaRule(
                formula=[f'AND(OR(${eur_below_letter}2=TRUE,${usd_below_letter}2=TRUE),${allow_letter}2=TRUE)'],
                fill=allowed_fill,
            ),
        )

    # Hide technical columns while retaining them for export/audit.
    for name in HIDDEN_COLUMNS:
        letter = get_column_letter(column_index[name])
        ws.column_dimensions[letter].hidden = True

    widths = {
        "Country": 24,
        "ISO": 11,
        "PricingUnitIdUsed": 20,
        "Plan": 14,
        "Days": 9,
        "GB": 9,
        "Price_EUR": 13,
        "Price_USD": 13,
        "PromoCode": 12,
        "FinalPriceAfterPromo_EUR": 19,
        "FinalPriceAfterPromo_USD": 19,
        "CostFloor_EUR": 14,
        "CostFloor_USD": 14,
        "EUR_IsBelowCostFloor": 17,
        "USD_IsBelowCostFloor": 17,
        "AllowBelowCost": 16,
    }
    for name, width in widths.items():
        ws.column_dimensions[get_column_letter(column_index[name])].width = width

    ws.row_dimensions[1].height = 26
    ws.sheet_properties.pageSetUpPr.fitToPage = True

    # Keep formula cells unlocked conceptually by styling only; no sheet protection.
    wb.save(path)
    return path
